"""Independent QA of ICAAP P0 items 2, 3, 4, 8 and M10 — the Appendix II return.

Agent 11 (Test/QA), 2026-09-19. Uses the real service chain (approved macro
scenario -> enterprise-stress run -> Board-attested sign-off -> package ->
exports) through the existing fixture helpers; every assertion is this file's
own. Defects are pinned as ``xfail(strict=True)`` with a ``P0-QA-###`` id.
"""

from __future__ import annotations

import ast
import copy
import csv
import io
import re
import zipfile
from decimal import Decimal
from pathlib import Path
from typing import Any
from uuid import UUID

import pdfplumber
import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.deps import TenantContext
from app.domain.stress.credit_bottom_up import CRD_EXPOSURE_CLASSES
from app.models import Bank, RegulatoryPackage, RegulatoryParameter, User
from app.models.stress import MACRO_VARIABLES
from app.services.attestation import digests
from app.services.regulatory_reporting.exports import export_package
from app.services.regulatory_reporting.templates import (
    APPENDIX2_EXPOSURE_CLASS_LABELS,
    APPENDIX2_RISK_DRIVER_LABELS,
    build_rendered_return,
    get_template,
)
from app.services.regulatory_reporting.validation import run_validation_rules
from tests.fixtures.canonical_bank_fixture import (
    DEMO_ORG_ID,
    SAMPLE_BANK_ID,
    materialize_canonical_test_book,
)
from tests.services.test_icaap_stress_appendix2_report import (
    MAKER,
    _approved_scenario,
    _attested_signoff,
    _generate,
    _generate_sdi,
    _run_enterprise_stress,
    _seed_checker,
    storage,
)
from tests.storage.inmemory import InMemoryStorageClient

__all__ = ["storage"]

# The enterprise stress run this suite prepares requires scoped FX and IRRBB
# calculation authority (an unconditional IRRBB check landed on the run
# service), so it takes the same two fixtures as its sibling
# ``test_icaap_stress_appendix2_report.py``.
pytestmark = pytest.mark.usefixtures("fx_run_authority", "irrbb_run_authority")

BACKEND = Path(__file__).parents[3]
PILLAR2_FIELDS = (
    "pillar2_credit_concentration",
    "pillar2_irrbb",
    "pillar2_sovereign",
    "pillar2_country_and_fx",
    "pillar2_reputational",
    "pillar2_others",
)
PILLAR2_LABELS = (
    "Credit Concentration",
    "IRRBB",
    "Sovereign",
    "Country and FX",
    "Reputational",
    "Others",
)
UNASSESSED_COLUMNS = ("current", "base_y1", "base_y2", "base_y3")

HOSTILE_ATTESTER_ID = UUID("eeeeeeee-eeee-4eee-8eee-eeeeeeeeeeee")
HOSTILE_ATTESTER = TenantContext(
    organization_id=DEMO_ORG_ID, actor_user_id=HOSTILE_ATTESTER_ID, roles=("approver",)
)
LONG_WORD = "X" * 2500
HOSTILE = {
    "scenario_narrative": (
        "=cmd|' /C calc'!A0 <script>alert('x')</script> <b>bold?</b> <para> "
        '<font color="red">red</font> & &amp; &lt; naïve café “quoted” — Tëst'
    ),
    "assumptions_rationale": f"+SUM(1,2) {LONG_WORD} END-OF-LONG-WORD",
    "methodology_summary": "-2+3 then " + ("Long paragraph sentence. " * 600) + "SENTINEL-TAIL",
    "board_challenge": (
        '@HYPERLINK("http://example.invalid") Line1\nLine2\r\n\nParagraph 2 <unclosed'
    ),
    "credibility_rationale": "Credible; <img src=x onerror=alert(1)> end.",
}


# --- helpers --------------------------------------------------------------------


def _section(snapshot: dict[str, Any], code: str) -> dict[str, Any] | None:
    return next((s for s in snapshot["sections"] if s["code"] == code), None)


def _stored(db: Session, storage_client: InMemoryStorageClient, object_path: str) -> bytes:
    slug = db.scalar(select(Bank.storage_slug).where(Bank.id == SAMPLE_BANK_ID))
    assert slug
    for obj in storage_client.list(slug, "outputs"):
        if obj.location.object_path == object_path:
            return obj and storage_client.read(obj.location)[1].read()
    raise AssertionError(object_path)


def _export(
    db: Session, storage_client: InMemoryStorageClient, package: RegulatoryPackage, kind: str
) -> bytes:
    artifact = export_package(db, MAKER, package, kind)  # type: ignore[arg-type]
    return _stored(db, storage_client, artifact.object_path)


def _pdf_pages(payload: bytes) -> list[tuple[float, float, str, list[dict[str, Any]]]]:
    with pdfplumber.open(io.BytesIO(payload)) as document:
        return [
            (float(page.width), float(page.height), page.extract_text() or "", list(page.chars))
            for page in document.pages
        ]


def _pdf_text(payload: bytes) -> str:
    return " ".join(" ".join(text for _w, _h, text, _c in _pdf_pages(payload)).split())


def _csv_entries(payload: bytes) -> dict[str, list[list[str]]]:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        return {
            name: list(csv.reader(io.StringIO(archive.read(name).decode("utf-8"))))
            for name in archive.namelist()
        }


def _csv_section(
    entries: dict[str, list[list[str]]], code: str
) -> tuple[list[str], list[list[str]]]:
    name = next(name for name in entries if name.endswith(f"_{code}.csv"))
    rows = entries[name]
    # header row = first row after the "#..." preamble
    start = next(i for i, row in enumerate(rows) if row and not row[0].startswith("#"))
    return rows[start], rows[start + 1 :]


def _prepare(
    db: Session, *, checker: TenantContext | None = None, **narrative: Any
) -> RegulatoryPackage:
    materialize_canonical_test_book(db)
    _seed_checker(db)
    scenario_id = _approved_scenario(db)
    run_id = _run_enterprise_stress(db, scenario_id)
    kwargs = dict(narrative)
    if checker is not None:
        kwargs["checker"] = checker
    _attested_signoff(db, run_id, **kwargs)
    return _generate(db)


def _seed_hostile_attester(db: Session) -> None:
    if db.get(User, HOSTILE_ATTESTER_ID) is not None:
        return
    db.add(
        User(
            id=HOSTILE_ATTESTER_ID,
            organization_id=DEMO_ORG_ID,
            email="hostile-board@aequoros.example",
            display_name='=HYPERLINK("http://evil.invalid","Chair")',
            job_title="<b>Board</b> & Chair",
            role="approver",
        )
    )
    db.commit()


# --- P0-2: neutral headers + provenance -----------------------------------------------


def _string_constants(path: Path) -> list[tuple[int, str]]:
    """Every string literal in a module that is NOT a docstring."""
    tree = ast.parse(path.read_text())
    docstrings: set[int] = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Module | ast.ClassDef | ast.FunctionDef | ast.AsyncFunctionDef):
            body = node.body
            if body and isinstance(body[0], ast.Expr) and isinstance(body[0].value, ast.Constant):
                docstrings.add(id(body[0].value))
    return [
        (node.lineno, node.value)
        for node in ast.walk(tree)
        if isinstance(node, ast.Constant)
        and isinstance(node.value, str)
        and id(node) not in docstrings
    ]


@pytest.mark.parametrize("module", ["templates.py", "generation.py"])
def test_no_13_percent_literal_outside_docstrings(module: str) -> None:
    path = BACKEND / "app" / "services" / "regulatory_reporting" / module
    offenders = [
        (line, text) for line, text in _string_constants(path) if re.search(r"13\s*%", text)
    ]
    assert offenders == []


def test_every_capital_requirement_header_is_neutral() -> None:
    for template_id in (
        "bog-icaap-stress-appendix2-v1",
        "bog-sdi-stress-annual-v1",
        "bog-bsd2-capital-v1",
    ):
        template = get_template(template_id)
        assert template is not None, template_id
        for layout in template.sections:
            texts = [layout.sheet_title, layout.source_citation, *layout.notes]
            texts += [column.header for column in layout.columns]
            for text in texts:
                assert not re.search(r"\b1[0-3](\.\d+)?\s*%", text), (template_id, text)


def test_governed_floor_provenance_is_in_the_snapshot_and_every_export(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    package = _prepare(db_session)
    provenance = {e["param_code"]: e for e in package.snapshot["metadata"]["parameter_provenance"]}
    car = provenance["car_min"]["governed"]
    for field in ("param_code", "value", "source_citation", "effective_from", "unit"):
        assert car[field] not in (None, ""), field
    assert Decimal(str(car["value"])) == Decimal("13")
    assert "paid_up_min" in provenance

    expected = (
        "Minimum total capital ratio applied: 13%",
        "Governed parameter car_min = 13%",
        str(car["effective_from"]),
    )
    pdf_text = _pdf_text(_export(db_session, storage, package, "pdf"))
    for fragment in (*expected, "Minimum unimpaired paid-up capital applied"):
        assert fragment in pdf_text, fragment
    assert "¶71" in pdf_text

    workbook = load_workbook(io.BytesIO(_export(db_session, storage, package, "xlsx")))
    meta_values = [
        str(cell.value) for row in workbook["Return Metadata"].iter_rows() for cell in row
    ]
    car_note = next(v for v in meta_values if v.startswith("Minimum total capital ratio applied"))
    for fragment in expected:
        assert fragment in car_note
    assert car["source_citation"] in car_note

    entries = _csv_entries(_export(db_session, storage, package, "csv"))
    notes = [row[1] for row in entries["00_metadata.csv"] if row and row[0] == "report_note"]
    car_csv = next(n for n in notes if n.startswith("Minimum total capital ratio applied"))
    for fragment in expected:
        assert fragment in car_csv
    assert car["source_citation"] in car_csv
    assert any(n.startswith("Minimum unimpaired paid-up capital applied") for n in notes)


def test_a_governed_car_min_of_11_5_prints_11_5(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    """Audit §4.2 acceptance: a fixture with car_min=11.5 renders 11.5%."""
    materialize_canonical_test_book(db_session)
    (row,) = db_session.scalars(
        select(RegulatoryParameter).where(
            RegulatoryParameter.param_code == "car_min", RegulatoryParameter.scope_key == "bank"
        )
    ).all()
    row.value_numeric = Decimal("11.5")
    db_session.commit()
    _seed_checker(db_session)
    run_id = _run_enterprise_stress(db_session, _approved_scenario(db_session))
    _attested_signoff(db_session, run_id)
    package = _generate(db_session)
    text = _pdf_text(_export(db_session, storage, package, "pdf"))
    assert "Minimum total capital ratio applied: 11.5%" in text
    assert "Governed parameter car_min = 11.5%" in text
    assert "13% CAR" not in text and "(13%)" not in text


# --- P0-3: Table 5 ----------------------------------------------------------------------


def test_table5_snapshot_invariants(db_session: Session) -> None:
    package = _prepare(db_session)
    pillar2 = {r["code"]: r for r in _section(package.snapshot, "t5_pillar2")["rows"]}  # type: ignore[index]
    rwa = {r["code"]: r for r in _section(package.snapshot, "t5_rwa")["rows"]}  # type: ignore[index]
    assert list(pillar2) == list(rwa)
    assert len(pillar2) == 7
    for code, row in pillar2.items():
        values = [row[field] for field in PILLAR2_FIELDS]
        modelled = [Decimal(str(v)) for v in values if v is not None]
        if code in UNASSESSED_COLUMNS:
            assert values == [None] * 6, code
            assert row["pillar2_total"] is None, code
        if modelled:
            assert abs(Decimal(str(row["pillar2_total"])) - sum(modelled)) <= Decimal("0.001")
        else:
            assert row["pillar2_total"] is None, code
        # Both grids agree on the total.
        assert rwa[code]["pillar2_total"] == row["pillar2_total"], code
        # Total requirement = Pillar 1 + Pillar 2 (nothing added for "not modelled").
        p1 = Decimal(str(rwa[code]["pillar1_requirement"]))
        p2 = Decimal(str(row["pillar2_total"])) if row["pillar2_total"] is not None else Decimal(0)
        total = Decimal(str(rwa[code]["total_capital_requirement"]))
        assert abs(total - (p1 + p2)) <= Decimal("0.01"), code
        assert str(row["pillar2_total"]) != "0.000"


def test_table5_not_modelled_in_pdf_xlsx_and_csv_never_zero(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    package = _prepare(db_session)

    text = _pdf_text(_export(db_session, storage, package, "pdf"))
    for label in (*PILLAR2_LABELS, "Total Pillar 2 Capital Requirements"):
        assert label in text, label
    # 4 unassessed columns x (6 risks + total) in the Pillar 2 grid + 4 totals in the RWA grid.
    assert text.count("Not modelled") >= 4 * 7 + 4

    workbook = load_workbook(io.BytesIO(_export(db_session, storage, package, "xlsx")))
    for sheet_marker, fields in (
        ("Pillar 2", (*PILLAR2_FIELDS, "pillar2_total")),
        ("Evolution of RWA", ("pillar2_total",)),
    ):
        name = next(n for n in workbook.sheetnames if sheet_marker in n)
        template = get_template("bog-icaap-stress-appendix2-v1")
        assert template is not None
        section_code = "t5_pillar2" if "Pillar 2" in sheet_marker else "t5_rwa"
        layout = next(s for s in template.sections if s.section_code == section_code)
        header_for = {c.key: c.header for c in layout.columns}
        rows = [list(row) for row in workbook[name].iter_rows()]
        header_idx = next(i for i, r in enumerate(rows) if r and r[0].value == header_for["code"])
        headers = [c.value for c in rows[header_idx]]
        for row in rows[header_idx + 1 :]:
            if row[0].value not in UNASSESSED_COLUMNS:
                continue
            for field in fields:
                cell = row[headers.index(header_for[field])]
                assert cell.value == "Not modelled", (name, row[0].value, field, cell.value)

    entries = _csv_entries(_export(db_session, storage, package, "csv"))
    for code, fields in (
        ("t5_pillar2", (*PILLAR2_FIELDS, "pillar2_total")),
        ("t5_rwa", ("pillar2_total",)),
    ):
        header, body = _csv_section(entries, code)
        template = get_template("bog-icaap-stress-appendix2-v1")
        assert template is not None
        layout = next(s for s in template.sections if s.section_code == code)
        header_for = {c.key: c.header for c in layout.columns}
        for row in body:
            if not row or row[0] not in UNASSESSED_COLUMNS:
                continue
            for field in fields:
                value = row[header.index(header_for[field])]
                assert value == "Not modelled", (code, row[0], field, value)


# --- P0-4: narratives -----------------------------------------------------------------


def test_hostile_narratives_are_literal_in_the_pdf_and_do_not_break_layout(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    _seed_hostile_attester(db_session)
    package = _prepare(db_session, checker=HOSTILE_ATTESTER, **HOSTILE)
    first = _export(db_session, storage, package, "pdf")
    second = _export(db_session, storage, package, "pdf")
    assert first == second, "PDF output is not byte-deterministic"

    pages = _pdf_pages(first)
    text = _pdf_text(first)
    for fragment in (
        "<script>alert('x')</script>",
        "<b>bold?</b>",
        "<para>",
        '<font color="red">red</font>',
        "& &amp; &lt;",
        "naïve café “quoted” — Tëst",
        "END-OF-LONG-WORD",
        "SENTINEL-TAIL",
        "Line1",
        "Paragraph 2 <unclosed",
        "<img src=x onerror=alert(1)>",
        '=HYPERLINK("http://evil.invalid","Chair"), <b>Board</b> & Chair',
    ):
        assert fragment in text, fragment
    # The 2,500-character word is wrapped inside the page, never run off it.
    for width, height, _text, chars in pages:
        for char in chars:
            assert char["x0"] >= -0.5 and char["x1"] <= width + 0.5, char["text"]
            assert char["top"] >= -0.5 and char["bottom"] <= height + 0.5
    # Narrative pages are portrait; the tabular pages before them stay landscape.
    narrative_start = next(
        i
        for i, (_w, _h, t, _c) in enumerate(pages)
        if "Stress Test Narrative (Board-attested)" in t
    )
    assert all(w < h for w, h, _t, _c in pages[narrative_start:])
    assert pages[narrative_start - 1][0] > pages[narrative_start - 1][1]
    assert len(pages) - narrative_start >= 2, "the 15k-character summary spills onto a new page"


@pytest.mark.xfail(
    strict=True,
    reason=(
        "P0-QA-001 (MEDIUM): the narrative PDF is set in the standard Helvetica font, which "
        "has no glyph for the cedi sign (U+20B5) or any non-WinAnsi character (CJK, emoji): "
        "reportlab prints a black-box placeholder, so a Board narrative citing 'GH₵' is not "
        "reproduced verbatim in the filed PDF. P0 fix round: the false 'verbatim' claim is "
        "removed, unprintable characters print as '?' and a note points to the signed "
        "snapshot (tests/services/test_icaap_p0_fix_round.py); the font itself is deferred "
        "to the P3 filing PDF (DV-004), so this stays pinned until then."
    ),
)
def test_the_cedi_sign_survives_into_the_narrative_pdf(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    package = _prepare(db_session, board_challenge="Capital shortfall of GH₵ 2.1bn challenged.")
    assert "GH₵ 2.1bn" in _pdf_text(_export(db_session, storage, package, "pdf"))


def test_hostile_narratives_are_inert_in_xlsx_and_csv(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    _seed_hostile_attester(db_session)
    package = _prepare(db_session, checker=HOSTILE_ATTESTER, **HOSTILE)
    expected = {
        key: value.replace("\r\n", "\n").replace("\r", "\n").strip()
        for key, value in HOSTILE.items()
    }
    rows = {r["code"]: r for r in _section(package.snapshot, "stress_narrative")["rows"]}  # type: ignore[index]
    for key, value in expected.items():
        assert rows[key]["value"] == value, key  # verbatim in the snapshot

    payload = _export(db_session, storage, package, "xlsx")
    workbook = load_workbook(io.BytesIO(payload))
    name = next(n for n in workbook.sheetnames if "Narrative" in n)
    cells = {str(row[0].value): row for row in workbook[name].iter_rows() if row and row[0].value}
    for key, value in expected.items():
        cell = cells[key][2]
        assert cell.value == value, key
        assert cell.data_type == "s", key
    assert str(cells["attested_by"][2].value).startswith('=HYPERLINK("http://evil.invalid"')
    assert cells["attested_by"][2].data_type == "s"
    # No formula element anywhere in the written sheet XML.
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        for entry in archive.namelist():
            if entry.startswith("xl/worksheets/"):
                assert "<f>" not in archive.read(entry).decode("utf-8"), entry

    entries = _csv_entries(_export(db_session, storage, package, "csv"))
    header, body = _csv_section(entries, "stress_narrative")
    by_key = {row[0]: row for row in body if row}
    statement = header.index("Statement")
    for key, value in expected.items():
        printed = by_key[key][statement]
        if value[0] in "=+-@":
            assert printed == "'" + value, key
        else:
            assert printed == value, key
    assert by_key["attested_by"][statement].startswith("'=HYPERLINK")
    # Every cell of the narrative CSV is inert.
    for row in body:
        for cell in row:
            assert not cell.startswith(("=", "+", "-", "@")), cell


def test_a_control_character_in_a_narrative_does_not_break_the_xlsx_export(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    package = _prepare(db_session, board_challenge="Soft\x0bline break from Word.")
    payload = _export(db_session, storage, package, "xlsx")
    assert payload


def test_control_character_narrative_still_renders_pdf_and_csv(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    package = _prepare(db_session, board_challenge="Soft\x0bline break from Word.")
    assert _export(db_session, storage, package, "pdf")
    assert _export(db_session, storage, package, "csv")


def test_narrative_change_moves_the_digest_and_identical_input_does_not(
    db_session: Session,
) -> None:
    package = _prepare(db_session)
    digest = digests.content_digest(package.snapshot)
    assert digests.content_digest(copy.deepcopy(package.snapshot)) == digest
    for index in range(6):
        altered = copy.deepcopy(package.snapshot)
        narrative = next(s for s in altered["sections"] if s["code"] == "stress_narrative")
        narrative["rows"][index]["value"] = "changed"
        assert digests.content_digest(altered) != digest, index


# --- M10 labels -------------------------------------------------------------------


def test_every_engine_key_has_a_printed_label() -> None:
    assert set(CRD_EXPOSURE_CLASSES) <= set(APPENDIX2_EXPOSURE_CLASS_LABELS)
    assert set(MACRO_VARIABLES) <= set(APPENDIX2_RISK_DRIVER_LABELS)
    for label in (
        *APPENDIX2_EXPOSURE_CLASS_LABELS.values(),
        *APPENDIX2_RISK_DRIVER_LABELS.values(),
    ):
        assert "_" not in label and label != label.upper(), label


def test_no_raw_enum_reaches_the_appendix_ii_pdf(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    package = _prepare(db_session)
    descriptions = {
        row["description"]
        for code in ("t1_impact_of_adverse", "t6_risk_drivers")
        for row in _section(package.snapshot, code)["rows"]  # type: ignore[index]
    }
    allowed = set(APPENDIX2_EXPOSURE_CLASS_LABELS.values()) | set(
        APPENDIX2_RISK_DRIVER_LABELS.values()
    )
    assert descriptions <= allowed
    text = _pdf_text(_export(db_session, storage, package, "pdf"))
    raw = {key.replace("_", " ").upper() for key in (*CRD_EXPOSURE_CLASSES, *MACRO_VARIABLES)}
    raw = {r for r in raw if " " in r} | {"GOG", "BOG"}
    leaked = sorted(r for r in raw if re.search(rf"\b{re.escape(r)}\b", text))
    assert leaked == []
    assert "FX rates (USD to GH Cedi)" in text
    assert "Government of Ghana" in text


# --- existing packages untouched ----------------------------------------------------


def test_exports_never_mutate_the_stored_snapshot_or_its_digest(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    package = _prepare(db_session)
    before = copy.deepcopy(package.snapshot)
    digest = package.content_digest
    for kind in ("pdf", "xlsx", "csv"):
        _export(db_session, storage, package, kind)
    db_session.refresh(package)
    assert package.snapshot == before
    assert package.content_digest == digest
    assert digests.content_digest(package.snapshot) == digests.content_digest(before)


def test_a_pre_p0_snapshot_renders_no_new_blocks_and_no_notes(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    """A snapshot shaped as generated before P0 (no t5_pillar2, no narrative, no
    provenance, a 0.000 Pillar 2 total) keeps its numbers: the zero is still a
    zero (null_label never rewrites a stored value) and no report note appears."""
    package = _prepare(db_session)
    legacy = copy.deepcopy(package.snapshot)
    legacy["sections"] = [
        s for s in legacy["sections"] if s["code"] not in {"t5_pillar2", "stress_narrative"}
    ]
    for key in ("parameter_provenance", "generation_findings", "report_notes"):
        legacy["metadata"].pop(key, None)
    for row in _section(legacy, "t5_rwa")["rows"]:  # type: ignore[index]
        if row["code"] in UNASSESSED_COLUMNS:
            row["pillar2_total"] = "0.000"
    rendered = build_rendered_return(
        get_template("bog-icaap-stress-appendix2-v1"),  # type: ignore[arg-type]
        legacy,
        package.source_runs,
        package_id=str(package.id),
        package_version=package.version,
    )
    assert rendered.report_notes == ()
    t5 = next(s for s in rendered.sections if s.layout.section_code == "t5_rwa")
    index = next(i for i, c in enumerate(t5.layout.columns) if c.key == "pillar2_total")
    assert all(row.cells[index].value == Decimal("0.000") for row in t5.rows[:4])


# --- SDI packet -------------------------------------------------------------------


def test_sdi_packet_renders_narrative_as_prose_with_its_own_floor(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    materialize_canonical_test_book(db_session)
    bank = db_session.get(Bank, SAMPLE_BANK_ID)
    assert bank is not None
    bank.institution_type = "savings_and_loans"
    db_session.flush()
    _seed_checker(db_session)
    run_id = _run_enterprise_stress(db_session, _approved_scenario(db_session, code="sdi_qa"))
    _attested_signoff(db_session, run_id)
    package = _generate_sdi(db_session)
    template = get_template("bog-sdi-stress-annual-v1")
    assert template is not None
    narrative = next(s for s in template.sections if s.section_code == "stress_narrative")
    assert narrative.presentation == "prose"
    text = _pdf_text(_export(db_session, storage, package, "pdf"))
    assert "Minimum total capital ratio applied: 10%" in text
    assert "Governed parameter car_min = 10%" in text


def test_sdi_pillar2_finding_matches_what_the_sdi_packet_prints(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    materialize_canonical_test_book(db_session)
    bank = db_session.get(Bank, SAMPLE_BANK_ID)
    assert bank is not None
    bank.institution_type = "savings_and_loans"
    db_session.flush()
    _seed_checker(db_session)
    run_id = _run_enterprise_stress(db_session, _approved_scenario(db_session, code="sdi_qa2"))
    _attested_signoff(db_session, run_id)
    package = _generate_sdi(db_session)
    findings = run_validation_rules(db_session, package)
    claims_not_modelled = any(
        f.get("rule") == "appendix2_pillar2_coverage" and "print 'Not modelled'" in f["detail"]
        for f in findings
    )
    rendered = build_rendered_return(
        get_template("bog-sdi-stress-annual-v1"),  # type: ignore[arg-type]
        package.snapshot,
        package.source_runs,
        package_id=str(package.id),
        package_version=package.version,
    )
    printed = {section.layout.section_code for section in rendered.sections}
    t5 = next(s for s in rendered.sections if s.layout.section_code == "t5_rwa")
    pillar2_columns = [c.key for c in t5.layout.columns if c.key.startswith("pillar2")]
    prints_pillar2 = "t5_pillar2" in printed or bool(pillar2_columns)
    # Fixed per decision D-020 (fix round 1): SDIs have no Pillar 2 regime, so
    # the packet neither raises the Pillar 2 finding nor prints a Pillar 2 grid
    # or column — the validation message and the artifact agree, in the
    # no-Pillar-2 direction (the original pin allowed either direction).
    assert claims_not_modelled == prints_pillar2
    assert not claims_not_modelled and not prints_pillar2
    assert "t5_pillar2" not in {s["code"] for s in package.snapshot["sections"]}


def test_a_pre_p0_package_still_states_the_minimum_it_was_measured_against(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    package = _prepare(db_session)
    legacy = copy.deepcopy(package.snapshot)
    legacy["sections"] = [
        s for s in legacy["sections"] if s["code"] not in {"t5_pillar2", "stress_narrative"}
    ]
    for key in ("parameter_provenance", "generation_findings", "report_notes"):
        legacy["metadata"].pop(key, None)
    assert Decimal(str(legacy["metadata"]["car_target_pct"])) == Decimal("13")
    rendered = build_rendered_return(
        get_template("bog-icaap-stress-appendix2-v1"),  # type: ignore[arg-type]
        legacy,
        package.source_runs,
        package_id=str(package.id),
        package_version=package.version,
    )
    printed = " ".join(
        [
            *(value for _label, value in rendered.metadata_pairs),
            *rendered.report_notes,
            *(s.layout.source_citation for s in rendered.sections),
            *(c.header for s in rendered.sections for c in s.layout.columns),
        ]
    )
    references_notes = "report notes" in printed
    states_minimum = "13%" in printed
    assert states_minimum or not references_notes
