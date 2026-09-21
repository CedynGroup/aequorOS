"""Independent QA of ICAAP P0 item 5 — reverse-stress summary in ICAAP-STRESS.

Agent 11 (Test/QA), 2026-09-19. Edge cases the implementation's own tests do
not cover: period scoping, failed runs, the latest-of-several rule, the run's
``input_hash`` binding, validation severity, every export format, and a
structural check that the STRESS-PACK builder is reused rather than copied.
"""

from __future__ import annotations

import ast
import csv
import io
import zipfile
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import BankReportingPeriod, RegulatoryPackage, RegulatoryRun
from app.schemas.reverse_stress import ReverseStressRunCreate
from app.services import reverse_stress
from app.services.regulatory_reporting import templates
from app.services.regulatory_reporting.exports import export_package
from app.services.regulatory_reporting.validation import run_validation_rules
from tests.fixtures.canonical_bank_fixture import DEMO_ORG_ID, SAMPLE_BANK_ID
from tests.services.test_icaap_stress_reverse_stress import (
    MAKER,
    _generate,
    _period_id,
    _seed,
    storage,
)
from tests.storage.inmemory import InMemoryStorageClient

__all__ = ["storage"]

GENERATION = Path(__file__).parents[3] / "app/services/regulatory_reporting/generation.py"


def _section(snapshot: dict[str, Any], code: str) -> dict[str, Any] | None:
    return next((s for s in snapshot["sections"] if s["code"] == code), None)


def _reverse_runs(db: Session) -> list[RegulatoryRun]:
    return list(
        db.scalars(
            select(RegulatoryRun)
            .where(
                RegulatoryRun.bank_id == SAMPLE_BANK_ID, RegulatoryRun.module == "reverse_stress"
            )
            .order_by(RegulatoryRun.created_at)
        ).all()
    )


def _stored(
    db: Session, client: InMemoryStorageClient, package: RegulatoryPackage, kind: str
) -> bytes:
    from app.models import Bank  # noqa: PLC0415

    artifact = export_package(db, MAKER, package, kind)  # type: ignore[arg-type]
    slug = db.scalar(select(Bank.storage_slug).where(Bank.id == SAMPLE_BANK_ID))
    for obj in client.list(slug, "outputs"):  # type: ignore[arg-type]
        if obj.location.object_path == artifact.object_path:
            return client.read(obj.location)[1].read()
    raise AssertionError(artifact.object_path)


def test_bound_run_carries_its_own_input_hash_and_appears_in_every_export(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    _seed(db_session, with_reverse=True)
    (run,) = _reverse_runs(db_session)
    package = _generate(db_session)
    (entry,) = [e for e in package.source_runs if e["module"] == "reverse_stress"]
    assert entry["run_id"] == str(run.id)
    assert entry["input_hash"] == run.input_hash and run.input_hash
    section = _section(package.snapshot, "reverse_stress")
    assert section is not None
    for row in section["rows"]:
        assert row["breached"] in {"true", "false"}
        assert row["scenario_code"]
        assert row["floor_pct"] is not None  # the governed floor used is stated
    findings = run_validation_rules(db_session, package)
    assert not [f for f in findings if f.get("rule") == "icaap_reverse_stress"]

    workbook = load_workbook(io.BytesIO(_stored(db_session, storage, package, "xlsx")))
    assert any("Reverse Stress" in name for name in workbook.sheetnames)
    with zipfile.ZipFile(io.BytesIO(_stored(db_session, storage, package, "csv"))) as archive:
        name = next(n for n in archive.namelist() if n.endswith("_reverse_stress.csv"))
        rows = list(csv.reader(io.StringIO(archive.read(name).decode("utf-8"))))
    codes = {row[0] for row in rows if row}
    assert {"liquidity_frontier", "capital_frontier"} <= codes


def test_a_reverse_run_for_another_period_is_not_bound(db_session: Session) -> None:
    _seed(db_session, with_reverse=True)
    (run,) = _reverse_runs(db_session)
    other = db_session.scalar(
        select(BankReportingPeriod.id).where(
            BankReportingPeriod.organization_id == DEMO_ORG_ID,
            BankReportingPeriod.bank_id == SAMPLE_BANK_ID,
            BankReportingPeriod.period_end == date(2025, 12, 31),
        )
    )
    assert other is not None and other != _period_id(db_session)
    run.reporting_period_id = other
    db_session.commit()
    package = _generate(db_session)
    assert _section(package.snapshot, "reverse_stress") is None
    assert all(e["module"] != "reverse_stress" for e in package.source_runs)
    assert package.snapshot["metadata"]["report_notes"][0].startswith("No reverse stress test")


def test_a_failed_reverse_run_is_never_bound(db_session: Session) -> None:
    _seed(db_session, with_reverse=True)
    (run,) = _reverse_runs(db_session)
    run.status = "failed"
    db_session.commit()
    package = _generate(db_session)
    assert _section(package.snapshot, "reverse_stress") is None
    findings = run_validation_rules(db_session, package)
    (info,) = [f for f in findings if f.get("rule") == "icaap_reverse_stress"]
    assert info["severity"] == "INFO"


def test_the_latest_of_several_runs_is_bound_by_both_packages(db_session: Session) -> None:
    _seed(db_session, with_reverse=True)
    reverse_stress.run_reverse_stress(
        db_session,
        MAKER,
        SAMPLE_BANK_ID,
        ReverseStressRunCreate(reporting_period_id=_period_id(db_session)),
    )
    runs = [r for r in _reverse_runs(db_session) if r.status == "succeeded"]
    assert len(runs) == 2
    latest = max(runs, key=lambda r: (r.created_at, str(r.id)))
    companion = _generate(db_session)
    pack = _generate(db_session, "STRESS-PACK")
    assert companion.snapshot["metadata"]["reverse_stress_run_id"] == str(latest.id)
    assert pack.snapshot["metadata"]["reverse_stress_run_id"] == str(latest.id)
    assert (
        _section(companion.snapshot, "reverse_stress")["rows"]  # type: ignore[index]
        == _section(pack.snapshot, "reverse_stress_frontier")["rows"]  # type: ignore[index]
    )


def test_absent_run_note_is_info_only_and_never_blocks(db_session: Session) -> None:
    _seed(db_session, with_reverse=False)
    package = _generate(db_session)
    findings = run_validation_rules(db_session, package)
    assert not [f for f in findings if f.get("severity") in {"ERROR", "BLOCKER"}]
    (info,) = [f for f in findings if f.get("rule") == "icaap_reverse_stress"]
    assert info["severity"] == "INFO"


def _calls(function: ast.FunctionDef) -> set[str]:
    return {
        node.func.id
        for node in ast.walk(function)
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name)
    }


def test_the_stress_pack_builder_is_reused_not_copied() -> None:
    tree = ast.parse(GENERATION.read_text())
    functions = {n.name: n for n in tree.body if isinstance(n, ast.FunctionDef)}
    for name in ("_generate_icaap_stress", "_generate_stress_pack"):
        calls = _calls(functions[name])
        assert "_latest_reverse_stress_run" in calls, name
    assert "_stress_frontier_rows" in _calls(functions["_generate_icaap_stress"])
    # One reverse-stress query and one frontier builder in the module.
    source = GENERATION.read_text()
    assert source.count('RegulatoryRun.module == "reverse_stress"') == 1
    assert source.count("lcr_at_breach_pct") == 1
    icaap = templates.get_template("bog-icaap-stress-v1")
    pack = templates.get_template("aeq-stress-pack-v1")
    assert icaap is not None and pack is not None
    icaap_cols = next(s for s in icaap.sections if s.section_code == "reverse_stress").columns
    pack_cols = next(
        s for s in pack.sections if s.section_code == "reverse_stress_frontier"
    ).columns
    assert icaap_cols is pack_cols


@pytest.mark.parametrize("with_reverse", [True, False])
def test_icaap_stress_pdf_is_byte_deterministic(
    db_session: Session, storage: InMemoryStorageClient, with_reverse: bool
) -> None:
    _seed(db_session, with_reverse=with_reverse)
    package = _generate(db_session)
    assert _stored(db_session, storage, package, "pdf") == _stored(
        db_session, storage, package, "pdf"
    )
