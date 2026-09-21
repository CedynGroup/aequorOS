"""ICAAP P0 fix round 1 (2026-09-19): regressions for the audit findings.

One test (or a small group) per coordinator item; the item number is in each
docstring. The Appendix II service chain is the real one (approved macro
scenario -> enterprise-stress run -> Board-attested sign-off -> package ->
exports), reusing the fixtures of ``test_icaap_stress_appendix2_report``.
"""

from __future__ import annotations

import ast
import io
import re
import zipfile
from dataclasses import replace
from datetime import date
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from typing import Any, cast
from uuid import UUID

import pdfplumber
import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.dialects import postgresql
from sqlalchemy.orm import Session

from app.models import Bank, RegulatoryPackage, RegulatoryRun
from app.services import capital_plan, enterprise_stress, regulatory_forecasting
from app.services.regulatory_reporting import generation, workflow
from app.services.regulatory_reporting.exports import export_package
from app.services.regulatory_reporting.exports.pdf_parts import escape_text, printable
from app.services.regulatory_reporting.exports.text import xml_safe
from app.services.regulatory_reporting.templates import (
    NOT_MODELLED,
    NOT_PROVIDED,
    build_rendered_return,
    get_template,
)
from app.services.regulatory_reporting.validation import run_validation_rules
from tests.fixtures.canonical_bank_fixture import SAMPLE_BANK_ID, materialize_canonical_test_book
from tests.fixtures.icaap.postgres_quarantine import NUL_IN_STORED_NARRATIVE
from tests.services.test_icaap_stress_appendix2_report import (
    MAKER,
    _approved_scenario,
    _attested_signoff,
    _generate,
    _generate_sdi,
    _run_enterprise_stress,
    _seed_checker,
    storage,
)
from tests.storage.inmemory import InMemoryStorageClient

__all__ = ["storage"]

# The enterprise stress run this suite prepares requires scoped FX and IRRBB
# calculation authority (an unconditional IRRBB check landed on the run
# service), so it takes the same two fixtures as the suite it borrows
# ``_run_enterprise_stress`` from.
pytestmark = pytest.mark.usefixtures("fx_run_authority", "irrbb_run_authority")

BACKEND = Path(__file__).parents[2]


# --- helpers ---------------------------------------------------------------------


def _section(snapshot: dict[str, Any], code: str) -> dict[str, Any] | None:
    return next((s for s in snapshot["sections"] if s["code"] == code), None)


def _export(
    db: Session, storage_client: InMemoryStorageClient, package: RegulatoryPackage, kind: str
) -> bytes:
    artifact = export_package(db, MAKER, package, kind)  # type: ignore[arg-type]
    slug = db.scalar(select(Bank.storage_slug).where(Bank.id == SAMPLE_BANK_ID))
    assert slug
    for obj in storage_client.list(slug, "outputs"):
        if obj.location.object_path == artifact.object_path:
            return storage_client.read(obj.location)[1].read()
    raise AssertionError(artifact.object_path)


def _pdf_text(payload: bytes) -> str:
    with pdfplumber.open(io.BytesIO(payload)) as document:
        text = "\n".join(page.extract_text() or "" for page in document.pages)
    return " ".join(text.split())


def _xlsx_strings(payload: bytes) -> list[str]:
    workbook = load_workbook(io.BytesIO(payload))
    return [
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]


def _csv_text(payload: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        return "\n".join(archive.read(name).decode("utf-8") for name in archive.namelist())


def _prepare(db: Session, **narrative: Any) -> RegulatoryPackage:
    materialize_canonical_test_book(db)
    _seed_checker(db)
    run_id = _run_enterprise_stress(db, _approved_scenario(db))
    _attested_signoff(db, run_id, **narrative)
    return _generate(db)


def _prepare_sdi(db: Session) -> RegulatoryPackage:
    materialize_canonical_test_book(db)
    bank = db.get(Bank, SAMPLE_BANK_ID)
    assert bank is not None
    bank.institution_type = "savings_and_loans"
    db.flush()
    _seed_checker(db)
    run_id = _run_enterprise_stress(db, _approved_scenario(db, code="sdi_fix_round"))
    _attested_signoff(db, run_id)
    return _generate_sdi(db)


def _findings(snapshot: dict[str, Any], rule: str) -> list[dict[str, Any]]:
    return [
        entry
        for entry in snapshot["metadata"].get("generation_findings", [])
        if entry.get("rule") == rule
    ]


def _provenance(package: RegulatoryPackage) -> dict[str, dict[str, Any]]:
    return {e["param_code"]: e for e in package.snapshot["metadata"]["parameter_provenance"]}


def _stress_run(db: Session, package: RegulatoryPackage) -> RegulatoryRun:
    (entry,) = package.source_runs
    run = db.get(RegulatoryRun, UUID(entry["run_id"]))
    assert run is not None
    return run


# --- item 1: SDIs have no Pillar 2 (D-020) ------------------------------------------


def test_the_sdi_packet_carries_no_pillar_2_anywhere(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    """Item 1: no Pillar 2 grid, column, note, citation, Table 5 reference or
    finding in the SDI snapshot, PDF, XLSX or CSV."""
    package = _prepare_sdi(db_session)
    snapshot = package.snapshot
    assert _section(snapshot, "t5_pillar2") is None
    rwa = _section(snapshot, "t5_rwa")
    assert rwa is not None
    for row in rwa["rows"]:
        assert not any(key.startswith("pillar2") for key in row), row
        assert "total_capital_requirement" not in row
    assert _findings(snapshot, "appendix2_pillar2_coverage") == []
    findings = run_validation_rules(db_session, package)
    assert not [f for f in findings if "pillar" in str(f.get("rule", "")).lower()]
    assert not [f for f in findings if "Pillar 2" in str(f.get("detail", ""))]

    forbidden = re.compile(r"Pillar[ -]?2|Table 5", re.IGNORECASE)
    pdf = _pdf_text(_export(db_session, storage, package, "pdf"))
    assert not forbidden.search(pdf), forbidden.search(pdf)
    for value in _xlsx_strings(_export(db_session, storage, package, "xlsx")):
        assert not forbidden.search(value), value
    csv_text = _csv_text(_export(db_session, storage, package, "csv"))
    assert not forbidden.search(csv_text)
    # The SDI packet's own requirement column is still there.
    assert "Capital Requirement at Governed CAR" in pdf


def test_an_sdi_stress_run_layers_no_pillar_2_add_on(db_session: Session) -> None:
    """Item 1 at the source: the SDI run's Table 5 requirement is the requirement
    at the governed CAR, with no engine-side add-on under another name."""
    package = _prepare_sdi(db_session)
    run = _stress_run(db_session, package)
    for row in run.metrics["appendix_ii"]["table5_rwa"]["rows"]:
        assert all(value is None for key, value in row["pillar2"].items() if key != "total")
        assert Decimal(row["total_capital_requirement"]) == Decimal(row["pillar1_requirement"])


# --- item 3: control characters never break an export ---------------------------------


@pytest.mark.parametrize(
    "control", ["\x0b", "\x0c", pytest.param("\x00", marks=NUL_IN_STORED_NARRATIVE)]
)
def test_a_control_character_in_a_stored_narrative_exports_to_every_format(
    db_session: Session, storage: InMemoryStorageClient, control: str
) -> None:
    """Item 3: a narrative stored with a control character (before the input
    validation existed) still exports to PDF, XLSX and CSV; the XLSX shows a
    line break for a vertical tab / form feed and a visible substitute otherwise."""
    package = _prepare(db_session, board_challenge=f"Line one{control}line two.")
    stored = _section(package.snapshot, "stress_narrative")
    assert stored is not None
    row = next(r for r in stored["rows"] if r["code"] == "board_challenge")
    assert row["value"] == f"Line one{control}line two."  # the snapshot keeps the exact text

    assert _export(db_session, storage, package, "pdf")
    strings = _xlsx_strings(_export(db_session, storage, package, "xlsx"))
    expected = "Line one\nline two." if control in "\x0b\x0c" else "Line one�line two."
    assert expected in strings
    assert _export(db_session, storage, package, "csv")


def test_submission_auto_mints_the_xlsx_despite_a_control_character(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    """Item 3: the filing set auto-exports what it is missing at submission; a
    stored control character must not turn that into a failure.

    The set is the FULL pack since 2026-09-20, so the XLSX the control
    character lives in is minted alongside the PDF record and the CSV. The
    point of the test is unchanged: none of them fails on the character."""
    package = _prepare(db_session, board_challenge="Soft\x0bline break from Word.")
    filed, detail = workflow._filing_set(db_session, MAKER, package)  # noqa: SLF001
    assert detail["auto_exported_kinds"] == ["xlsx", "pdf", "csv"]
    assert sorted(artifact.kind for artifact in filed) == ["csv", "pdf", "xlsx"]


def test_xml_safe_is_the_identity_on_ordinary_text() -> None:
    text = "Tab\there, line\nbreak, carriage\rreturn, naïve café — ¶67(c)"
    assert xml_safe(text) is text
    assert xml_safe("a\x0bb\x0cc\x00d\x1fe") == "a\nb\nc�d�e"


# --- item 5: a weaker applied minimum is never called stricter -------------------------


def test_a_weaker_applied_minimum_is_labelled_weaker_and_warned(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    """Item 5: applied below the governed value -> ``weaker_than_governed``, a
    WARNING, and truthful note wording; equal -> ``governed``; above ->
    ``stricter_than_governed``."""
    materialize_canonical_test_book(db_session)
    _seed_checker(db_session)
    run_id = _run_enterprise_stress(db_session, _approved_scenario(db_session))
    run = db_session.get(RegulatoryRun, run_id)
    assert run is not None
    # A run sealed before the clamp adoption could record the governed 15% row
    # while applying a lower value: stage exactly that.
    ledger = [dict(entry) for entry in run.parameter_provenance or []]
    for entry in ledger:
        if entry.get("param_code") == "car_min":
            entry["value"] = "15.000000"
    run.parameter_provenance = ledger
    db_session.commit()
    _attested_signoff(db_session, run_id)
    package = _generate(db_session)

    car = _provenance(package)["car_min"]
    assert car["basis"] == "weaker_than_governed"
    (warning,) = _findings(package.snapshot, "appendix2_minimum_weaker_than_governed")
    assert warning["severity"] == "WARNING"
    assert "understate" in warning["detail"]
    text = _pdf_text(_export(db_session, storage, package, "pdf"))
    assert "is LOWER than the governed minimum" in text
    assert "stricter than the governed minimum" not in text


def test_the_basis_is_three_way() -> None:
    governed = {"value": "13"}
    assert generation._minimum_basis("13", governed, "percent") == "governed"  # noqa: SLF001
    assert generation._minimum_basis("15", governed, "percent") == "stricter_than_governed"  # noqa: SLF001
    assert generation._minimum_basis("10", governed, "percent") == "weaker_than_governed"  # noqa: SLF001
    assert generation._minimum_basis("10", None, "percent") == "not_recorded"  # noqa: SLF001


# --- item 6: ¶67 citations follow the directive ------------------------------------


def test_methodologies_are_cited_as_67c_never_67d() -> None:
    """Item 6: ¶67(c) is the methodologies, ¶67(d) the impact on profitability,
    capital and liquidity — nothing cites ¶67(d) for methodologies."""
    template = get_template("bog-icaap-stress-appendix2-v1")
    assert template is not None
    narrative = next(s for s in template.sections if s.section_code == "stress_narrative")
    assert re.search(r"¶67\(c\)[^;]*methodolog", narrative.source_citation)
    labels = dict(generation._APPENDIX2_NARRATIVE_ELEMENTS)  # noqa: SLF001
    assert "¶67(c)" in labels["methodology_summary"]
    assert "¶67(b)" in labels["assumptions_rationale"]
    wrong = re.compile(r"67\(d\)[^.\n]{0,40}methodolog", re.IGNORECASE)
    for path in (
        BACKEND / "app" / "models" / "stress.py",
        BACKEND / "app" / "schemas" / "enterprise_stress_signoff.py",
        BACKEND / "app" / "services" / "regulatory_reporting" / "templates.py",
        BACKEND / "app" / "services" / "regulatory_reporting" / "generation.py",
        BACKEND / "dashboard" / "components" / "stress" / "SignoffPanel.tsx",
    ):
        assert not wrong.search(path.read_text(encoding="utf-8")), path


# --- item 8: nil is not "not modelled" ---------------------------------------------


def test_an_assessed_nil_charge_is_zero_not_absent() -> None:
    """Item 8 at the source: a module that ran and found no charge records 0.

    The FX arm now reads the add-on the shared ICAAP function produced
    (``pillar2_addon``, D-038) instead of re-deriving a NOP ratio here; an FX
    book that GAINS under the scenario is the assessed nil.
    """
    outcome = SimpleNamespace(
        concentration=SimpleNamespace(pillar2_concentration_charge=Decimal("0")),
        irr=SimpleNamespace(delta_eve=Decimal("5000")),  # an EVE gain: no loss
        fx=SimpleNamespace(pillar2_addon=Decimal("0")),  # a long book, cedi falls
        operational=None,
    )
    overlay = enterprise_stress._pillar2_overlay(cast(Any, outcome), 3)  # noqa: SLF001
    assert overlay is not None
    year1 = overlay[1]
    assert year1.credit_concentration == Decimal("0")
    assert year1.irrbb == Decimal("0")
    assert year1.country_and_fx == Decimal("0")
    assert year1.other is None  # the operational module did not run: not modelled
    nothing_ran = SimpleNamespace(concentration=None, irr=None, fx=None, operational=None)
    assert enterprise_stress._pillar2_overlay(cast(Any, nothing_ran), 3) is None  # noqa: SLF001


def test_table5_prints_a_nil_charge_as_zero_and_an_absent_one_as_not_modelled() -> None:
    row = {
        "label": "stress_y1",
        "pillar2": {
            "credit_concentration": "0.000",
            "irrbb": "120.000",
            "sovereign": None,
            "country_and_fx": "0.000",
            "reputational": None,
            "other": "30.000",
            "total": "150.000",
        },
    }
    pillar2, missing = generation._appendix2_pillar2(row)  # noqa: SLF001
    assert missing == ["Sovereign", "Reputational"]
    assert pillar2["pillar2_credit_concentration"] == "0.000"
    assert pillar2["pillar2_total"] == "150.000"
    assert pillar2["pillar2_coverage"] == (
        "Partial — excludes Sovereign, Reputational; assessed nil: Credit Concentration, "
        "Country and FX"
    )
    full = get_template("bog-icaap-stress-appendix2-v1")
    assert full is not None
    template = replace(
        full, sections=tuple(s for s in full.sections if s.section_code == "t5_pillar2")
    )
    snapshot = {
        "template_id": template.template_id,
        "metadata": {"template_revision": 2, "parameter_provenance": []},
        "sections": [
            {
                "code": "t5_pillar2",
                "rows": [{"code": "stress_y1", "description": "Stress", **pillar2}],
            }
        ],
    }
    rendered = build_rendered_return(template, snapshot, [], package_id="p", package_version=1)
    grid = next(s for s in rendered.sections if s.layout.section_code == "t5_pillar2")
    cells = {
        spec.key: cell for spec, cell in zip(grid.layout.columns, grid.rows[0].cells, strict=True)
    }
    assert cells["pillar2_credit_concentration"].value == Decimal("0.000")
    assert cells["pillar2_sovereign"].value == NOT_MODELLED
    notes = " ".join(grid.layout.notes)
    assert "assessed at nil prints 0.00" in notes


# --- item 9: the capital-return citation makes no retroactive claim -----------------


def test_the_capital_return_citation_is_conditional_on_the_run_date() -> None:
    template = get_template("bog-bsd2-capital-v1")
    assert template is not None
    citation = next(
        s for s in template.sections if s.section_code == "capital_ratios"
    ).source_citation
    assert "19 September 2026" in citation
    assert "a run computed before that date governed the CAR only" in citation
    assert "may only tighten" in citation


# --- item 10: every applied minimum is stated -------------------------------------


def test_the_report_notes_state_every_applied_minimum(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    """Item 10: CAR, CET1, Tier 1 and leverage (and paid-up) with code, value and
    citation; a run that applied an ungoverned 3% leverage says so."""
    materialize_canonical_test_book(db_session)
    _seed_checker(db_session)
    run_id = _run_enterprise_stress(db_session, _approved_scenario(db_session))
    run = db_session.get(RegulatoryRun, run_id)
    assert run is not None
    # Stage a pre-B2 run: leverage applied at the old 3% with no governed row.
    inputs = dict(run.inputs)
    parameters = dict(inputs["parameters"])
    capital = dict(parameters["capital"])
    thresholds = dict(capital["thresholds_pct"])
    thresholds["leverage_min"] = "3"
    capital["thresholds_pct"] = thresholds
    parameters["capital"] = capital
    inputs["parameters"] = parameters
    run.inputs = inputs
    run.parameter_provenance = [
        e for e in (run.parameter_provenance or []) if e.get("param_code") != "leverage_min"
    ]
    db_session.commit()
    _attested_signoff(db_session, run_id)
    package = _generate(db_session)

    provenance = _provenance(package)
    assert list(provenance) == [
        "car_min",
        "cet1_min",
        "tier1_min",
        "leverage_min",
        "paid_up_min",
        # Fix round 2 (D-024 / M21): the governed Table 2 recognition caps.
        "at1_cap_pct_rwa",
        "tier2_cap_pct_rwa",
    ]
    assert provenance["cet1_min"]["basis"] == "governed"
    assert provenance["leverage_min"]["basis"] == "not_recorded"
    text = _pdf_text(_export(db_session, storage, package, "pdf"))
    for label, code, value in (
        ("Minimum CET1 ratio", "cet1_min", "6.5%"),
        ("Minimum Tier 1 ratio", "tier1_min", "8%"),
    ):
        assert f"{label} applied: {value}" in text
        assert f"Governed parameter {code} = {value}" in text
        assert provenance[code]["governed"]["source_citation"] in text
    assert "Minimum leverage ratio applied: 3%. The source run does not record" in text


# --- item 12: characters the standard font cannot draw ------------------------------


def test_an_undrawable_character_prints_as_a_visible_substitute(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    package = _prepare(db_session, board_challenge="Shortfall of GH₵ 2.1bn challenged.")
    text = _pdf_text(_export(db_session, storage, package, "pdf"))
    assert "Shortfall of GH? 2.1bn challenged." in text
    assert "The signed package snapshot, the XLSX and the CSV hold the exact text" in text
    assert "reproduced without editing" not in text and "verbatim from the Board" not in text
    assert "Shortfall of GH₵ 2.1bn challenged." in _xlsx_strings(
        _export(db_session, storage, package, "xlsx")
    )


def test_printable_keeps_the_windows_1252_repertoire() -> None:
    kept = "naïve café “quoted” — Tëst ¶67 €5"
    assert printable(kept) == kept
    assert printable("₵ 日本 😀") == "? ?? ?"
    assert escape_text("<b>a</b> & ₵") == "&lt;b&gt;a&lt;/b&gt; &amp; ?"


# --- item 15: the governed row the resolver applied ---------------------------------


def test_the_governed_row_follows_the_resolver_layer_order() -> None:
    """Item 15: a licence-type row wins over the class row whatever their dates."""
    ledger = [
        {
            "param_code": "car_min",
            "scope_type": "institution_class",
            "scope_key": "bank",
            "value": "13",
            "effective_from": "2026-01-01",
            "effective_to": None,
            "parameter_id": "b",
        },
        {
            "param_code": "car_min",
            "scope_type": "institution_type",
            "scope_key": "universal_bank",
            "value": "14",
            "effective_from": "2020-01-01",
            "effective_to": None,
            "parameter_id": "a",
        },
    ]
    run = cast(RegulatoryRun, SimpleNamespace(parameter_provenance=ledger))
    row = generation._governed_row(run, "car_min", date(2026, 3, 31))  # noqa: SLF001
    assert row is not None and row["scope_type"] == "institution_type"
    class_only = cast(RegulatoryRun, SimpleNamespace(parameter_provenance=ledger[:1]))
    row = generation._governed_row(class_only, "car_min", date(2026, 3, 31))  # noqa: SLF001
    assert row is not None and row["value"] == "13"


# --- item 17: one horizon rule --------------------------------------------------------


def test_both_icaap_consumers_share_one_horizon_rule() -> None:
    compiled = {
        str(select(RegulatoryRun.id).where(clause).compile(dialect=postgresql.dialect()))
        for clause in (
            capital_plan._planning_horizon(),  # noqa: SLF001
            regulatory_forecasting.regulatory_horizon_clause(),
        )
    }
    assert len(compiled) == 1
    source = (BACKEND / "app" / "services" / "regulatory_reporting" / "generation.py").read_text(
        encoding="utf-8"
    )
    tree = ast.parse(source)
    function = next(
        node
        for node in tree.body
        if isinstance(node, ast.FunctionDef) and node.name == "_generate_icaap_stress"
    )
    body = ast.unparse(function)
    assert "regulatory_horizon_clause()" in body
    assert "horizon_years" not in body


# --- item 18: printed wording ----------------------------------------------------------


def test_the_gse_driver_prints_the_year_on_year_change_never_the_level(
    db_session: Session, storage: InMemoryStorageClient
) -> None:
    package = _prepare(db_session)
    rows = [
        r
        for r in _section(package.snapshot, "t6_risk_drivers")["rows"]  # type: ignore[index]
        if r["code"].startswith("gse_index:")
    ]
    assert rows
    by_year = {r["year_index"]: r for r in rows}
    assert by_year["1"]["base_value"] is None and by_year["1"]["stress_value"] is None
    for year in ("2", "3"):
        # The fixture holds the index flat, so the change is nil — never the level.
        assert Decimal(by_year[year]["base_value"]) == 0
        assert Decimal(by_year[year]["stress_value"]) == 0
    assert all(
        r["description"] == "Year-on-Year Changes in Stock Market Valuation (GSE Index)"
        for r in rows
    )
    text = _pdf_text(_export(db_session, storage, package, "pdf"))
    assert NOT_PROVIDED in text
    # The authored levels (5,000 base / 3,500 stress) never reach the page.
    assert not re.search(r"(?<![\d,])(5,000|3,500)\.00", text)


def test_year_on_year_is_computed_from_consecutive_levels() -> None:
    section = generation._appendix2_table6_section(  # noqa: SLF001
        {
            "rows": [
                {
                    "variable": "gse_index",
                    "year_index": 1,
                    "base_value": "5000",
                    "stress_value": "4000",
                },
                {
                    "variable": "gse_index",
                    "year_index": 2,
                    "base_value": "5500",
                    "stress_value": "3000",
                },
                {
                    "variable": "gdp_growth",
                    "year_index": 1,
                    "base_value": "0.05",
                    "stress_value": "0",
                },
            ]
        }
    )
    rows = {row["code"]: row for row in section["rows"]}
    assert rows["gse_index:y1"]["base_value"] is None
    assert Decimal(rows["gse_index:y2"]["base_value"]) == Decimal("0.1")
    assert Decimal(rows["gse_index:y2"]["stress_value"]) == Decimal("-0.25")
    assert rows["gdp_growth:y1"]["base_value"] == "0.05"  # other drivers untouched


def test_governance_rows_are_in_words(db_session: Session) -> None:
    package = _prepare(db_session)
    governance = _section(package.snapshot, "governance")
    assert governance is not None
    rows = {row["code"]: row for row in governance["rows"]}
    assert "attested_by" not in rows  # the name is in the narrative, no user id here
    assert rows["signoff_status"]["value"] == "Attested by the Board"


def test_bog_template_headers_use_the_directive_wording() -> None:
    template = get_template("bog-icaap-stress-appendix2-v1")
    assert template is not None
    headers = {column.header for section in template.sections for column in section.columns}
    for expected in (
        "RWA for Credit Risk",
        "RWA for Operational Risk",
        "RWA for Market Risk",
        "Total Pillar 1 RWA",
        "Pillar 1 Capital Requirements",
        "Capital Required to Meet the Minimum Unimpaired Paid-up Capital",
    ):
        assert expected in headers, expected
    t5 = next(s for s in template.sections if s.section_code == "t5_rwa")
    keys = [column.key for column in t5.columns]
    # Components first, then the total, as the directive prints them.
    assert keys.index("credit_rwa") < keys.index("value")


def test_the_irr_note_names_the_discounting_curve() -> None:
    template = get_template("bog-irrbb-pilot-v1")
    assert template is not None
    eve = next(s for s in template.sections if s.section_code == "eve_scenarios")
    notes = " ".join(eve.notes)
    assert "single projection curve (discounting on the published OIS curve" in notes


def test_the_reverse_stress_citation_states_its_proxy_and_its_pre_action_basis() -> None:
    template = get_template("bog-icaap-stress-v1")
    assert template is not None
    section = next(s for s in template.sections if s.section_code == "reverse_stress")
    assert "insolvency or illiquidity" in section.source_citation
    assert "proxy" in section.source_citation
    assert "before management actions" in " ".join(section.notes)
    assert any(column.key == "scenario" for column in section.columns)
