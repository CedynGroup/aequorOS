"""The ICAAP working copy (DOCX).

The DOCX is the copy a bank's reviewers redline, so the risks are different
from the PDF's: it is editable and therefore must never be mistaken for the
filed instrument, and it is a zip container whose writer stamps wall-clock
times into every entry, which would make two exports of the same document
differ. Both are checked here, along with the structure a reviewer relies on —
real Word headings, real tables, and no style the default template does not
ship.

Unlike the PDF, OOXML carries UTF-8: the cedi sign and Ghanaian vowels must
survive exactly, with no substitution.
"""

from __future__ import annotations

import io
import zipfile
from datetime import date, datetime

import pytest
from docx import Document as read_document

from app.domain.icaap.document import (
    DRAFT_WATERMARK,
    NOT_AVAILABLE,
    REHEARSAL_WATERMARK,
    WORKING_COPY_NOTICE,
    AttachmentLine,
    ChecklistLine,
    DataBlockRef,
    FactRef,
    IcaapDocument,
    ListBlock,
    ParagraphBlock,
    QuoteBlock,
    Run,
    SectionRender,
)
from app.services.icaap.render import format as fmt
from app.services.icaap.render.docx import render_docx

GENERATED_AT = datetime(2026, 9, 19, 12, 0, 0)

#: Everything the python-docx default template ships that this renderer uses.
#: A style outside this set means the document opens with a missing style in
#: Word, which shows as unformatted text.
EXPECTED_STYLES = {
    "Normal",
    "Heading 1",
    "Heading 2",
    "Heading 3",
    "Heading 4",
    "Quote",
    "List Paragraph",
}


def paragraph(text: str, style: str = "body") -> ParagraphBlock:
    return ParagraphBlock(runs=(Run(text=text),), style=style)  # type: ignore[arg-type]


def opened(payload: bytes):  # noqa: ANN201 - python-docx has no exported Document type
    return read_document(io.BytesIO(payload))


def body_text(payload: bytes) -> str:
    document = opened(payload)
    lines = [paragraph_.text for paragraph_ in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.extend(cell.text for cell in row.cells)
    return "\n".join(lines)


def capital_block() -> object:
    payload = {
        "schema": fmt.PAYLOAD_SCHEMA,
        "title": "Capital position",
        "as_of": "2025-12-31",
        "source_label": "Official capital run (baseline)",
        "unit": {"currency": "GHS", "scale": 1},
        "tables": [
            {
                "key": "capital_components",
                "title": "Capital components",
                "columns": [
                    {"key": "label", "label": "Component", "kind": "text"},
                    {"key": "amount", "label": "Amount", "kind": "amount"},
                ],
                "rows": [
                    {"cells": {"label": "CET1 capital", "amount": "1234567"}},
                    {"cells": {"label": "Total capital", "amount": "-25"}, "emphasis": "total"},
                ],
            }
        ],
    }
    return fmt.block_render(
        block_id="b1",
        payload=payload,
        status="fresh",
        fallback_title="Capital position",
        provenance=(("Run", "run-abc"),),
    )


def build_document(
    *,
    sections: tuple[SectionRender, ...] | None = None,
    is_rehearsal: bool = False,
    blocks: dict[str, object] | None = None,
    fact_text: dict[tuple[str, str], str] | None = None,
    attachments: tuple[AttachmentLine, ...] = (),
) -> IcaapDocument:
    if sections is None:
        sections = (
            SectionRender(
                key="executive_summary",
                letter="a",
                title="Executive Summary",
                citation_label="¶49(a)",
                source_status="sourced",
                content_label="Version 2 (committed)",
                blocks=(
                    paragraph("Risk appetite", "h2"),
                    paragraph("The Board reviewed the capital position."),
                ),
                checklist=(
                    ChecklistLine(
                        item_id="a1",
                        text="State the purpose of the ICAAP",
                        citation_label="¶49(a)",
                        status="not_applicable",
                        reason="The group has no subsidiaries",
                    ),
                ),
            ),
        )
    return IcaapDocument(
        institution_name="Sample Bank Ltd",
        institution_short_name="Sample Bank",
        cycle_title="ICAAP FY2025",
        fiscal_year=2025,
        as_of=date(2025, 12, 31),
        basis="solo",
        basis_label="Solo",
        cycle_kind="rehearsal" if is_rehearsal else "annual",
        cycle_kind_label="Rehearsal" if is_rehearsal else "Annual ICAAP",
        framework_code="BOG-ICAAP",
        framework_title="Guideline on ICAAP",
        framework_version="2026.02-ed",
        framework_status="exposure_draft",
        framework_digest="sha256:7f3c",
        regulator_short="BoG",
        currency="GHS",
        generated_at=GENERATED_AT,
        sections=sections,
        blocks=blocks or {},  # type: ignore[arg-type]
        fact_text=fact_text or {},
        attachments=attachments,
        readiness_summary=("2 items block freeze",),
        is_rehearsal=is_rehearsal,
    )


class TestDeterminism:
    def test_two_renders_are_byte_identical(self) -> None:
        # python-docx stamps wall-clock times into every zip entry at save, so
        # this only holds because the container is re-normalised.
        document = build_document(blocks={"b1": capital_block()})
        assert render_docx(document) == render_docx(document)

    def test_every_zip_entry_carries_the_fixed_timestamp(self) -> None:
        with zipfile.ZipFile(io.BytesIO(render_docx(build_document()))) as archive:
            assert archive.namelist() == sorted(archive.namelist())
            for info in archive.infolist():
                assert info.date_time == (1980, 1, 1, 0, 0, 0)

    def test_the_document_properties_are_pinned_to_the_caller(self) -> None:
        properties = opened(render_docx(build_document())).core_properties
        # OOXML stores these as UTC instants; python-docx reads them back aware.
        assert properties.created is not None
        assert properties.modified is not None
        assert properties.created.replace(tzinfo=None) == GENERATED_AT
        assert properties.modified.replace(tzinfo=None) == GENERATED_AT
        assert properties.author == "AequorOS"
        assert properties.revision == 1

    def test_a_re_render_produces_the_same_text(self) -> None:
        document = build_document(blocks={"b1": capital_block()})
        assert body_text(render_docx(document)) == body_text(render_docx(document))


class TestWorkingCopyNotice:
    def test_the_header_says_it_is_not_the_filed_document(self) -> None:
        document = opened(render_docx(build_document()))
        headers = [section.header.paragraphs[0].text for section in document.sections]
        assert headers
        for header in headers:
            assert WORKING_COPY_NOTICE in header
            assert DRAFT_WATERMARK in header

    def test_a_rehearsal_header_says_so_too(self) -> None:
        document = opened(render_docx(build_document(is_rehearsal=True)))
        for section in document.sections:
            assert REHEARSAL_WATERMARK in section.header.paragraphs[0].text

    def test_a_normal_cycle_does_not_claim_to_be_a_rehearsal(self) -> None:
        document = opened(render_docx(build_document()))
        for section in document.sections:
            assert REHEARSAL_WATERMARK not in section.header.paragraphs[0].text

    def test_the_notice_is_repeated_in_the_body_and_the_properties(self) -> None:
        payload = render_docx(build_document())
        assert WORKING_COPY_NOTICE in body_text(payload)
        assert opened(payload).core_properties.comments == WORKING_COPY_NOTICE


class TestStructure:
    def test_sections_are_real_word_headings(self) -> None:
        document = opened(render_docx(build_document()))
        headings = [
            paragraph_.text
            for paragraph_ in document.paragraphs
            if paragraph_.style is not None and paragraph_.style.name == "Heading 1"
        ]
        assert "(a) Executive Summary" in headings
        assert "Contents" in headings
        assert "Requirement checklist" in headings
        assert "Provenance" in headings

    def test_editor_headings_keep_their_level(self) -> None:
        document = opened(render_docx(build_document()))
        levels = {
            paragraph_.text: paragraph_.style.name
            for paragraph_ in document.paragraphs
            if paragraph_.style is not None
        }
        assert levels["Risk appetite"] == "Heading 2"

    def test_no_style_outside_the_default_template_is_used(self) -> None:
        document = opened(render_docx(build_document(blocks={"b1": capital_block()})))
        used = {
            paragraph_.style.name
            for paragraph_ in document.paragraphs
            if paragraph_.style is not None
        }
        assert used <= EXPECTED_STYLES
        for table in document.tables:
            assert table.style is not None
            assert table.style.name == "Table Grid"

    def test_figures_are_real_tables_not_text(self) -> None:
        section = SectionRender(
            key="s",
            letter="a",
            title="Section",
            citation_label="",
            source_status="sourced",
            content_label="Working draft — not committed",
            blocks=(DataBlockRef(block_id="b1"),),
        )
        document = opened(
            render_docx(build_document(sections=(section,), blocks={"b1": capital_block()}))
        )
        grids = [
            table
            for table in document.tables
            if table.rows and table.rows[0].cells[0].text == "Component"
        ]
        assert len(grids) == 1
        assert grids[0].rows[1].cells[1].text == "1,234,567.00"

    def test_marks_survive_as_word_formatting(self) -> None:
        section = SectionRender(
            key="s",
            letter="a",
            title="Section",
            citation_label="",
            source_status="sourced",
            content_label="Working draft — not committed",
            blocks=(
                ParagraphBlock(
                    runs=(
                        Run(text="bold", bold=True),
                        Run(text="italic", italic=True),
                        Run(text="under", underline=True),
                    ),
                    style="body",
                ),
            ),
        )
        document = opened(render_docx(build_document(sections=(section,))))
        runs = {
            run.text: (run.bold, run.italic, run.underline)
            for paragraph_ in document.paragraphs
            for run in paragraph_.runs
        }
        assert runs["bold"][0] is True
        assert runs["italic"][1] is True
        assert runs["under"][2] is True

    def test_lists_carry_explicit_markers_that_honour_the_start(self) -> None:
        # Word numbering definitions cannot be restarted per list
        # deterministically, so the markers are written out.
        section = SectionRender(
            key="s",
            letter="a",
            title="Section",
            citation_label="",
            source_status="sourced",
            content_label="Working draft — not committed",
            blocks=(
                ListBlock(
                    ordered=True,
                    start=3,
                    items=((paragraph("third item"),), (paragraph("fourth item"),)),
                ),
                ListBlock(ordered=False, start=1, items=((paragraph("a bullet"),),)),
                QuoteBlock(children=(paragraph("a quoted line"),)),
            ),
        )
        document = opened(render_docx(build_document(sections=(section,))))
        texts = [paragraph_.text for paragraph_ in document.paragraphs]
        assert "3. third item" in texts
        assert "4. fourth item" in texts
        assert "• a bullet" in texts
        quotes = [
            paragraph_.text
            for paragraph_ in document.paragraphs
            if paragraph_.style is not None and paragraph_.style.name == "Quote"
        ]
        assert quotes == ["a quoted line"]


class TestContent:
    def test_unicode_survives_exactly(self) -> None:
        # The PDF substitutes these (DV-004); OOXML has no such limit, so the
        # DOCX is the copy that holds the author's characters.
        section = SectionRender(
            key="s",
            letter="a",
            title="Section",
            citation_label="",
            source_status="sourced",
            content_label="Working draft — not committed",
            blocks=(paragraph("Capital of ₵1m, reviewed by Ɛ. Ɔsei."),),
        )
        assert "₵1m" in body_text(render_docx(build_document(sections=(section,))))
        assert "Ɔsei" in body_text(render_docx(build_document(sections=(section,))))

    @pytest.mark.parametrize(
        "hostile",
        [
            "<b>bold</b> & <script>alert(1)</script>",
            "</w:t></w:r>",
            "a < b > c & d",
        ],
    )
    def test_markup_prints_as_characters(self, hostile: str) -> None:
        section = SectionRender(
            key="s",
            letter="a",
            title="Section",
            citation_label="",
            source_status="sourced",
            content_label="Working draft — not committed",
            blocks=(paragraph(hostile),),
        )
        assert hostile in body_text(render_docx(build_document(sections=(section,))))

    def test_control_characters_do_not_break_the_export(self) -> None:
        section = SectionRender(
            key="s",
            letter="a",
            title="Section",
            citation_label="",
            source_status="sourced",
            content_label="Working draft — not committed",
            blocks=(paragraph("before\x0bafter\x07bell\x00nul"),),
        )
        text = body_text(render_docx(build_document(sections=(section,))))
        assert "before" in text
        assert "after" in text

    def test_a_very_long_word_still_produces_a_document(self) -> None:
        section = SectionRender(
            key="s",
            letter="a",
            title="Section",
            citation_label="",
            source_status="sourced",
            content_label="Working draft — not committed",
            blocks=(paragraph("A" * 4000),),
        )
        assert "A" * 4000 in body_text(render_docx(build_document(sections=(section,))))

    def test_a_factref_prints_the_bound_figure(self) -> None:
        section = SectionRender(
            key="s",
            letter="a",
            title="Section",
            citation_label="",
            source_status="sourced",
            content_label="Working draft — not committed",
            blocks=(
                ParagraphBlock(
                    runs=(
                        Run(text="CAR is "),
                        Run(text="", fact=FactRef(block_id="b1", fact_key="car_pct")),
                    ),
                    style="body",
                ),
            ),
        )
        text = body_text(
            render_docx(
                build_document(sections=(section,), fact_text={("b1", "car_pct"): "15.42%"})
            )
        )
        assert "CAR is 15.42%" in text

    def test_a_fact_with_no_value_prints_not_available_not_zero(self) -> None:
        section = SectionRender(
            key="s",
            letter="a",
            title="Section",
            citation_label="",
            source_status="sourced",
            content_label="Working draft — not committed",
            blocks=(
                ParagraphBlock(
                    runs=(
                        Run(text="Buffer: "),
                        Run(text="", fact=FactRef(block_id="b1", fact_key="absent")),
                    ),
                    style="body",
                ),
            ),
        )
        assert f"Buffer: {NOT_AVAILABLE}" in body_text(
            render_docx(build_document(sections=(section,)))
        )

    def test_a_referenced_block_the_cycle_no_longer_has_is_named(self) -> None:
        section = SectionRender(
            key="s",
            letter="a",
            title="Section",
            citation_label="",
            source_status="sourced",
            content_label="Working draft — not committed",
            blocks=(DataBlockRef(block_id="gone"),),
        )
        assert "no longer part of this cycle" in body_text(
            render_docx(build_document(sections=(section,)))
        )

    def test_an_empty_section_prints_only_its_own_note(self) -> None:
        section = SectionRender(
            key="s",
            letter="q",
            title="Disclosure",
            citation_label="¶82",
            source_status="sourced",
            content_label="Working draft — not committed",
            blocks=(),
            empty_note="No text has been written for this section yet.",
        )
        assert "No text has been written for this section yet." in body_text(
            render_docx(build_document(sections=(section,)))
        )

    def test_a_cycle_with_no_sections_still_renders(self) -> None:
        text = body_text(render_docx(build_document(sections=())))
        assert "This cycle has no sections." in text
        assert "No requirements are recorded" in text

    def test_the_reason_a_requirement_is_not_applicable_reaches_the_docx(self) -> None:
        assert "The group has no subsidiaries" in body_text(render_docx(build_document()))

    def test_attachments_are_indexed_with_their_checksum(self) -> None:
        attachments = (
            AttachmentLine(
                kind_title="Board resolution",
                title="Resolution 12/2026",
                sha256_prefix="9f2a1c04",
                uploaded_on=date(2026, 1, 5),
            ),
        )
        assert "9f2a1c04" in body_text(render_docx(build_document(attachments=attachments)))

    def test_a_framework_that_moved_under_the_cycle_is_declared(self) -> None:
        document = build_document()
        object.__setattr__(
            document,
            "framework_digest_note",
            "This cycle recorded framework digest sha256:old; this export renders sha256:7f3c.",
        )
        assert "recorded framework digest sha256:old" in body_text(render_docx(document))

    def test_the_provenance_names_the_framework_digest(self) -> None:
        assert "sha256:7f3c" in body_text(render_docx(build_document()))
