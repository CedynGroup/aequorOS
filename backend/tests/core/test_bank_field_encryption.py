from __future__ import annotations

import json
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import cast
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import TypeAdapter
from sqlalchemy import Uuid, bindparam, select, text
from sqlalchemy.orm import Session, undefer, undefer_group

from app.attestation.service import register_state
from app.attestation.service.digests import register_state_digest
from app.behavioral.ml import deposit_stability, nmd_duration
from app.behavioral.ml.config import BehavioralTrainingConfig
from app.behavioral.ml.history import load_deposit_month_aggregates
from app.bi.domain.catalogue import catalogue
from app.bi.models.content import BiDashboard, BiDashboardVersion
from app.bi.models.marts import BiDimCounterparty, BiFactPositionDaily, BiFactPositionEom
from app.bi.models.notifications import BiAlert, BiSubscription
from app.bi.schemas.content import BiDashboardSpec
from app.bi.schemas.query import BiFilter, BiPackQuery, BiQuery, BiTime
from app.bi.service import content
from app.bi.service.compiler import compile_query
from app.bi.service.errors import BiQueryError
from app.bi.service.execution import execute
from app.bi.service.identity_policy import IDENTITY_REASON
from app.core.ids import new_uuid7
from app.core.key_management import registry
from app.core.key_management.concentration_identity import counterparty_group_key
from app.core.key_management.fields import Normalization, blind_index, observe_bank_key
from app.core.key_management.local import LocalKeyProvider
from app.core.key_management.models import ObjectKeyEnvelope
from app.core.key_management.presentation import (
    LOCKED_JSON,
    LOCKED_TEXT,
    SESSION_STATE,
    SESSION_SUPPORT,
    PresentationState,
)
from app.core.key_management.types import (
    KeyIntegrityError,
    KeyReference,
    KeyStatus,
    KeyUnavailableError,
)
from app.core.tenancy import TenantContext
from app.credit.service import credit_concentration, credit_exposure_book
from app.db.base import Base, utc_now
from app.domain.ingestion.contracts import EntityMapping, MappingConfig
from app.identity.models.institution_profile import RelatedParty, RelatedPartyRole, Shareholding
from app.identity.public import Bank, User
from app.identity.schemas.institution_profile import (
    InstitutionProfileFullRead,
    RelatedPartyCreate,
    RelatedPartyRoleInput,
    RelatedPartyUpdate,
)
from app.identity.service import banks, institution_profile
from app.live import position_book
from app.live.service import fact_derivation
from app.models import ParamConcentrationLimit
from app.models.audit_event import AuditEvent
from app.models.canonical import CanonicalCounterparty, CanonicalPosition, CanonicalPositionSnapshot
from app.models.canonical_withdrawal import CanonicalWithdrawal
from app.models.ingestion import IngestionBatch, LineageRecord, TranslationFailure
from app.reporting.models.packages import RegulatoryPackage
from app.schemas.ingestion import IngestionBatchCreate, MappingConfigCreate
from app.services import canonical_withdrawal, ingestion
from app.services.behavioral_liquidity import get_behavioral_liquidity_report
from app.services.canonical_withdrawal import WithdrawalError
from tests.api import test_impersonation_auth as impersonation
from tests.bi.fixtures import materialize_bi_plane
from tests.support.helpers import ORG_1, ORG_2, USER_1, headers
from tests.support.inmemory_storage import InMemoryStorageClient

BANK_A = "BK-SAMP0001"
BANK_B = "BK-SAMP0002"
AS_OF = date(2026, 6, 30)
KEY_A = KeyReference("aws_kms", "synthetic-bank-a", "us-east-1", "111122223333")
KEY_B = KeyReference("aws_kms", "synthetic-bank-b", "us-east-1", "111122223333")
NEXT = KeyReference("aws_kms", "synthetic-next-a", "us-east-1", "111122223333")
NAME = "Synthetic Borrower Ñ"
EXPANDING_NAME = "\ufb03" * 130
IDENTIFIER = "SYNTHETIC-CARD-17"
CTX = TenantContext(organization_id=ORG_1, actor_user_id=USER_1, roles=("admin",))


@pytest.fixture
def provider(db_session: Session, monkeypatch: pytest.MonkeyPatch) -> LocalKeyProvider:
    provider = LocalKeyProvider()
    for key in (KEY_A, KEY_B, NEXT):
        provider.add_key(key)

    def resolve(_name: str) -> LocalKeyProvider:
        return provider

    monkeypatch.setattr(registry, "provider_for", resolve)
    for bank_id in (BANK_A, BANK_B):
        db_session.add(
            Bank(
                id=bank_id,
                organization_id=ORG_1,
                name="Synthetic test bank",
                short_name="Synthetic",
                jurisdiction_code="GH",
                currency="GHS",
                license_type="universal_bank",
                institution_type="universal_bank",
                storage_slug=bank_id.lower(),
            )
        )
    db_session.flush()
    return provider


def connect(db: Session, provider: LocalKeyProvider, bank_id: str = BANK_A) -> None:
    registry.register(
        db,
        bank_id=bank_id,
        organization_id=ORG_1,
        storage_slug=bank_id.lower(),
        key=KEY_A if bank_id == BANK_A else KEY_B,
        providers=lambda _name: provider,
    )
    db.commit()


def counterparty(  # noqa: PLR0913
    db: Session,
    bank_id: str = BANK_A,
    reference: str = IDENTIFIER,
    *,
    include_package: bool = True,
    source_system: str = "EXCEL_CSV",
    as_of: date = AS_OF,
    name: str = NAME,
) -> CanonicalCounterparty:
    batch = IngestionBatch(
        organization_id=ORG_1,
        bank_id=bank_id,
        source_system=source_system,
        adapter_version="synthetic",
        extraction_mode="full",
        status="accepted",
        as_of_date=as_of,
        validation_report={"name": NAME},
        etl_report={"name": NAME},
        error_message=NAME,
    )
    db.add(batch)
    db.flush()
    lineage = LineageRecord(
        organization_id=ORG_1,
        bank_id=bank_id,
        ingestion_batch_id=batch.id,
        operation_type="ADAPTER_TRANSLATE",
        operation_ref="synthetic-test",
        input_lineage_ids=[],
        details={"name": NAME},
    )
    db.add(lineage)
    db.flush()
    row = CanonicalCounterparty(
        organization_id=ORG_1,
        bank_id=bank_id,
        as_of_date=as_of,
        source_system=source_system,
        source_reference=reference,
        ingestion_batch_id=batch.id,
        lineage_id=lineage.id,
        validation_status="accepted",
        name=name,
        counterparty_type="RETAIL_INDIVIDUAL",
        group_reference="SYNTHETIC-GROUP-17",
        external_identifiers={"national_id": IDENTIFIER},
        attributes={"contact_name": NAME},
    )
    db.add(row)
    db.add(
        TranslationFailure(
            organization_id=ORG_1,
            bank_id=bank_id,
            ingestion_batch_id=batch.id,
            entity_type="counterparty",
            source_locator="synthetic-row-1",
            raw_record={"name": NAME, "national_id": IDENTIFIER},
            error_code="synthetic",
            error_message=NAME,
        )
    )
    if include_package:
        db.add(
            RegulatoryPackage(
                organization_id=ORG_1,
                bank_id=bank_id,
                return_family="large_exposures",
                return_code=reference,
                reporting_date=as_of,
                frequency="monthly",
                status="generated",
                version=1,
                snapshot={"counterparty_name": NAME, "balance": "100"},
                validation_report={"message": NAME},
                notes=NAME,
                generated_by=USER_1,
            )
        )
    db.flush()
    return row


def raw_fields(db: Session) -> dict[str, list[tuple[object, ...]]]:
    # This explicit catalogue is the coverage contract, rather than a reflection
    # of whichever columns happen to have an encryption implementation today.
    catalogue = {
        "canonical_counterparties": (
            "name, external_identifiers, group_reference, source_reference"
        ),
        "related_parties": "full_name, contact",
        "related_party_roles": "other_responsibilities, term_of_appointment, icag_registration",
        "ingestion_batches": "validation_report, etl_report, error_message",
        "lineage_records": "details",
        "translation_failures": "raw_record, error_message",
    }
    return {
        table: [
            tuple(row) for row in db.execute(text(f"SELECT {columns} FROM {table} ORDER BY id"))
        ]
        for table, columns in catalogue.items()
    }


def test_every_covered_field_is_ciphertext_and_entity_reads_round_trip(
    db_session: Session,
    provider: LocalKeyProvider,
) -> None:
    connect(db_session, provider)
    cp = counterparty(db_session)
    cp_id = cp.id
    party = RelatedParty(
        organization_id=ORG_1,
        bank_id=BANK_A,
        party_type="individual",
        full_name=NAME,
        contact={"email": "synthetic@example.test", "national_id": IDENTIFIER},
    )
    db_session.add(party)
    db_session.flush()
    db_session.add(
        RelatedPartyRole(
            organization_id=ORG_1,
            party_id=party.id,
            role="director",
            other_responsibilities=NAME,
            term_of_appointment=NAME,
            icag_registration=IDENTIFIER,
            sitting_allowance=Decimal("12.34"),
        )
    )
    db_session.add(
        Shareholding(
            organization_id=ORG_1,
            party_id=party.id,
            share_type="ordinary",
            shareholder_rights="voting",
            number_of_shares=Decimal("100"),
            pct_shareholding=Decimal("25"),
            ubo_party_id=party.id,
        )
    )
    db_session.commit()
    stored = raw_fields(db_session)
    assert all(rows for rows in stored.values())
    for table, rows in stored.items():
        for row in rows:
            for value in row:
                encoded = json.dumps(value, ensure_ascii=False, default=str)
                assert "aqs-bank-field-v1" in encoded, (table, encoded)
                assert NAME not in encoded
                assert IDENTIFIER not in encoded
                assert "synthetic@example.test" not in encoded
    allowance = cast(
        object,
        db_session.execute(text("SELECT sitting_allowance FROM related_party_roles")).scalar_one(),
    )
    assert Decimal(str(allowance)) == Decimal("12.34")
    assert (
        db_session.execute(text("SELECT number_of_shares FROM shareholdings")).scalar_one() == 100
    )
    db_session.expire_all()
    loaded = db_session.get(CanonicalCounterparty, cp_id)
    assert loaded is not None and loaded.name == NAME and loaded.source_reference == IDENTIFIER
    assert loaded.external_identifiers == {"national_id": IDENTIFIER}
    loaded_party = db_session.scalar(select(RelatedParty))
    assert loaded_party is not None and loaded_party.full_name == NAME
    assert loaded_party.contact["national_id"] == IDENTIFIER
    loaded_role = db_session.scalar(select(RelatedPartyRole))
    assert loaded_role is not None and loaded_role.icag_registration == IDENTIFIER
    assert loaded_role.other_responsibilities == NAME and loaded_role.term_of_appointment == NAME
    batch = db_session.scalar(select(IngestionBatch))
    lineage = db_session.scalar(select(LineageRecord))
    failure = db_session.scalar(select(TranslationFailure))
    assert batch is not None and batch.validation_report == {"name": NAME}
    assert lineage is not None and lineage.details == {"name": NAME}
    assert failure is not None and failure.raw_record["name"] == NAME
    package = db_session.scalar(select(RegulatoryPackage))
    assert package is not None and package.snapshot["counterparty_name"] == NAME
    assert package.validation_report == {"message": NAME} and package.notes == NAME


def test_bank_scoped_indexes_and_master_key_rotation(
    db_session: Session,
    provider: LocalKeyProvider,
) -> None:
    connect(db_session, provider)
    connect(db_session, provider, BANK_B)
    a = counterparty(db_session)
    counterparty(db_session, BANK_B)
    a_id = a.id
    db_session.commit()
    before = raw_fields(db_session)
    indexes = cast(
        list[object],
        list(
            db_session.execute(
                text("SELECT name_blind_index FROM canonical_counterparties ORDER BY bank_id")
            ).scalars()
        ),
    )
    assert len(set(indexes)) == 2
    assert all(NAME not in str(token) and len(str(token)) == 64 for token in indexes)
    tokens: tuple[tuple[str, str, Normalization, object], ...] = (
        ("counterparty.name", "  SYNTHETIC   borrower ñ  ", "name", a.name_blind_index),
        ("counterparty.source_reference", IDENTIFIER, "identifier", a.source_reference_blind_index),
        (
            "counterparty.identifier:national_id",
            IDENTIFIER,
            "identifier",
            a.external_identifiers_blind_indexes["national_id"],
        ),
    )
    for domain, value, normalization, expected in tokens:
        assert (
            blind_index(
                db_session,
                organization_id=ORG_1,
                bank_id=BANK_A,
                domain=domain,
                value=value,
                normalization=normalization,
            )
            == expected
        )
    assert (
        blind_index(
            db_session,
            organization_id=ORG_1,
            bank_id=BANK_A,
            domain="counterparty.source_reference",
            value=IDENTIFIER.lower(),
        )
        != a.source_reference_blind_index
    )
    assert (
        blind_index(
            db_session,
            organization_id=ORG_1,
            bank_id=BANK_A,
            domain="counterparty.identifier:tax_id",
            value=IDENTIFIER,
        )
        != a.external_identifiers_blind_indexes["national_id"]
    )
    row = registry.scoped_key(db_session, bank_id=BANK_A, organization_id=ORG_1)
    registry.rotate(db_session, row, NEXT, lambda _provider: provider)
    db_session.commit()
    db_session.expunge_all()
    loaded = db_session.get(CanonicalCounterparty, a_id)
    assert loaded is not None and loaded.name == NAME
    assert raw_fields(db_session) == before


@pytest.mark.parametrize("substitution", ["bank", "record", "field"])
def test_ciphertext_cannot_be_moved_between_banks_records_or_fields(
    db_session: Session,
    provider: LocalKeyProvider,
    substitution: str,
) -> None:
    connect(db_session, provider)
    connect(db_session, provider, BANK_B)
    a = counterparty(db_session)
    b = counterparty(db_session, BANK_B if substitution == "bank" else BANK_A, "SYNTHETIC-OTHER")
    a_id, b_id = a.id, b.id
    db_session.commit()
    column = "group_reference" if substitution == "field" else "name"
    source = cast(
        object,
        db_session.execute(
            text(f"SELECT {column} FROM canonical_counterparties WHERE id=:id"),
            {"id": b_id if db_session.get_bind().dialect.name == "postgresql" else b_id.hex},
        ).scalar_one(),
    )
    assert isinstance(source, str)
    db_session.execute(
        text("UPDATE canonical_counterparties SET name=:value WHERE id=:id"),
        {
            "id": a_id if db_session.get_bind().dialect.name == "postgresql" else a_id.hex,
            "value": source,
        },
    )
    db_session.commit()
    db_session.expunge_all()
    with pytest.raises(KeyUnavailableError):
        loaded = db_session.get(CanonicalCounterparty, a_id)
        assert loaded is not None
        _ = loaded.name


def test_bank_a_key_cannot_unwrap_bank_b_and_revocation_refuses_reads(
    db_session: Session,
    provider: LocalKeyProvider,
) -> None:
    connect(db_session, provider)
    connect(db_session, provider, BANK_B)
    b_id = counterparty(db_session, BANK_B).id
    db_session.commit()
    envelope = db_session.scalar(
        select(ObjectKeyEnvelope).where(ObjectKeyEnvelope.bank_id == BANK_B)
    )
    assert envelope is not None
    with pytest.raises(KeyUnavailableError):
        provider.unwrap(
            KEY_A, envelope.wrapped_key, registry.wrapping_context(BANK_B, envelope.context_digest)
        )
    provider.set_status(KEY_B, KeyStatus.DISABLED)
    db_session.expunge_all()
    with pytest.raises(KeyUnavailableError):
        loaded = db_session.get(CanonicalCounterparty, b_id)
        assert loaded is not None
        _ = loaded.name


def test_register_crud_and_new_audits_omit_names(
    db_session: Session,
    provider: LocalKeyProvider,
) -> None:
    connect(db_session, provider)
    payload = RelatedPartyCreate(
        reason="Synthetic privacy test",
        party_type="individual",
        full_name=NAME,
        contact={"national_id": IDENTIFIER},
        roles=[
            RelatedPartyRoleInput(role="director"),
            RelatedPartyRoleInput(role="ultimate_beneficial_owner"),
        ],
    )
    created = institution_profile.create_related_party(db_session, CTX, BANK_A, payload)
    assert created.full_name == NAME
    db_session.expire_all()
    profile = institution_profile.get_full_profile(db_session, CTX, BANK_A)
    assert profile.related_parties[0].contact == {"national_id": IDENTIFIER}
    audits = list(
        db_session.scalars(
            select(AuditEvent).where(AuditEvent.event_type == "related_party.created")
        )
    )
    assert len(audits) == 1 and "full_name" not in audits[0].details
    assert NAME not in json.dumps(audits[0].details, ensure_ascii=False)
    updated = institution_profile.update_related_party(
        db_session,
        CTX,
        BANK_A,
        created.id,
        RelatedPartyUpdate.model_validate(
            {**payload.model_dump(), "full_name": "Synthetic replacement"}
        ),
    )
    assert updated.full_name == "Synthetic replacement"
    db_session.expunge_all()
    profile = institution_profile.get_full_profile(db_session, CTX, BANK_A)
    assert profile.related_parties[0].full_name == "Synthetic replacement"
    event = db_session.scalar(
        select(AuditEvent).where(AuditEvent.event_type == "related_party.updated")
    )
    assert event is not None and "full_name" not in event.details


def assert_unicode_name_round_trip(db: Session, connected: bool) -> None:
    borrower = counterparty(db, include_package=False, name=EXPANDING_NAME)
    borrower_id = borrower.id
    party = institution_profile.create_related_party(
        db,
        CTX,
        BANK_A,
        RelatedPartyCreate(
            reason="Synthetic Unicode names", party_type="individual", full_name=EXPANDING_NAME
        ),
    )
    for name, equivalent in ((EXPANDING_NAME, "FFI" * 130), ("\ufb04" * 130, "FFL" * 130)):
        if name != EXPANDING_NAME:
            borrower = db.get(CanonicalCounterparty, borrower_id)
            assert borrower is not None
            borrower.name = name
            institution_profile.update_related_party(
                db,
                CTX,
                BANK_A,
                party.id,
                RelatedPartyUpdate(
                    reason="Synthetic Unicode rename", party_type="individual", full_name=name
                ),
            )
        db.expunge_all()
        loaded = db.get(CanonicalCounterparty, borrower_id)
        assert loaded is not None and loaded.name == name
        profile = institution_profile.get_full_profile(db, CTX, BANK_A)
        assert profile.related_parties[0].full_name == name
        stored = db.execute(
            text(
                "SELECT name, name_blind_index FROM canonical_counterparties "
                "UNION ALL SELECT full_name, full_name_blind_index FROM related_parties"
            )
        ).all()
        assert len(stored) == 2
        for record, domain in zip(
            stored, ("counterparty.name", "related_party.full_name"), strict=True
        ):
            stored_name, token = cast(object, record[0]), cast(object, record[1])
            assert token == blind_index(
                db,
                organization_id=ORG_1,
                bank_id=BANK_A,
                domain=domain,
                value=equivalent,
                normalization="name",
            )
            if connected:
                assert stored_name != name
            else:
                assert stored_name == name


@pytest.mark.parametrize("connected", [False, True])
def test_unicode_name_writes_updates_and_normalized_indexes(
    db_session: Session, provider: LocalKeyProvider, connected: bool
) -> None:
    if connected:
        connect(db_session, provider)
    assert_unicode_name_round_trip(db_session, connected)


def test_numeric_sql_aggregation_preserves_opaque_concentration_groups(
    db_session: Session,
    provider: LocalKeyProvider,
) -> None:
    connect(db_session, provider)
    first = counterparty(db_session)
    second = counterparty(db_session, reference="SYNTHETIC-SECOND")
    second.name = "Synthetic second borrower"
    for n, party in enumerate((first, second), start=1):
        common = dict(
            organization_id=ORG_1,
            bank_id=BANK_A,
            as_of_date=AS_OF,
            source_system="EXCEL_CSV",
            ingestion_batch_id=party.ingestion_batch_id,
            lineage_id=party.lineage_id,
            validation_status="accepted",
        )
        position = CanonicalPosition(
            **common,
            source_reference=f"synthetic-deposit-{n}",
            position_type="DEPOSIT",
            currency="GHS",
        )
        db_session.add(position)
        db_session.flush()
        db_session.add(
            CanonicalPositionSnapshot(
                **common,
                source_reference=f"synthetic-snapshot-{n}",
                position_id=position.id,
                counterparty_id=party.id,
                balance=Decimal(100 * n),
            )
        )
    db_session.commit()
    provider.set_status(KEY_A, KeyStatus.DISABLED)
    db_session.expunge_all()
    aggregates = load_deposit_month_aggregates(
        db_session, CTX, BANK_A, AS_OF, 1, non_maturing_only=False
    )
    assert len(aggregates) == 1
    assert len(aggregates[0].concentration_group) == 64
    assert NAME not in aggregates[0].concentration_group
    assert "SYNTHETIC-GROUP-17" not in aggregates[0].concentration_group
    assert aggregates[0].balance_ghs == 300 and aggregates[0].n_accounts == 2


def test_partial_refresh_and_update_preserve_other_encrypted_fields(
    db_session: Session,
    provider: LocalKeyProvider,
) -> None:
    connect(db_session, provider)
    row_id = counterparty(db_session).id
    db_session.commit()
    db_session.expunge_all()
    row = db_session.get(CanonicalCounterparty, row_id)
    assert row is not None
    db_session.refresh(row, ["name"])
    assert row.name == NAME
    db_session.expire(row, ["attributes", "external_identifiers", "group_reference"])
    row.name = "Synthetic replacement borrower"
    db_session.commit()
    db_session.expunge_all()
    loaded = db_session.get(CanonicalCounterparty, row_id)
    assert loaded is not None and loaded.name == "Synthetic replacement borrower"
    assert loaded.source_reference == IDENTIFIER
    assert loaded.attributes == {"contact_name": NAME}
    assert loaded.group_reference == "SYNTHETIC-GROUP-17"
    assert loaded.external_identifiers == {"national_id": IDENTIFIER}


def test_failed_sensitive_write_rolls_back_without_a_platform_fallback(
    db_session: Session,
    provider: LocalKeyProvider,
) -> None:
    connect(db_session, provider)
    row = counterparty(db_session)
    row_id = row.id
    db_session.commit()
    before = raw_fields(db_session)
    provider.set_status(KEY_A, KeyStatus.DISABLED)
    row.name = "Synthetic refused write"
    with pytest.raises(KeyUnavailableError):
        db_session.flush()
    db_session.rollback()
    provider.set_status(KEY_A, KeyStatus.ACTIVE)
    loaded = db_session.get(CanonicalCounterparty, row_id)
    assert loaded is not None and loaded.name == NAME
    assert raw_fields(db_session) == before


def withdraw_counterparties(db: Session, *, expected_rows: int = 1) -> CanonicalWithdrawal:
    checker = User(
        id=new_uuid7(),
        organization_id=ORG_1,
        email=f"synthetic-checker-{new_uuid7()}@example.test",
        display_name="Synthetic checker",
    )
    db.add(checker)
    db.flush()
    bank = db.get(Bank, BANK_A)
    assert bank is not None
    requested = canonical_withdrawal.request_withdrawal(
        db,
        CTX,
        bank,
        entity="counterparty",
        source_system="EXCEL_CSV",
        as_of_date=AS_OF,
        reason="Synthetic book withdrawal",
        requested_by="synthetic-analyst@example.test",
    )
    applied = canonical_withdrawal.approve_withdrawal(
        db,
        TenantContext(organization_id=ORG_1, actor_user_id=checker.id),
        bank,
        requested.id,
        approved_by=checker.email,
    )
    assert applied.status == "applied" and applied.rows_withdrawn == expected_rows
    db.commit()
    return applied


def write_counterparty_workbook(path: Path, name: str) -> None:
    workbook = Workbook()
    customers = workbook.active
    assert customers is not None
    customers.title = "Customers"
    customers.append(["CustomerId", "CustomerName", "Type"])
    customers.append([IDENTIFIER, name, "RETAIL_INDIVIDUAL"])
    loans = cast(Worksheet, workbook.create_sheet("Loans"))
    loans.append(["Reference", "Type", "Currency", "Balance", "CustomerId"])
    loans.append(["synthetic-loan-1", "LOAN", "GHS", 123, IDENTIFIER])
    workbook.save(path)


@pytest.mark.parametrize(
    "connection_timing,withdraw",
    [
        ("none", False),
        ("none", True),
        ("before", False),
        ("before", True),
        ("after_withdrawal", True),
    ],
)
@pytest.mark.parametrize("initial_name", [NAME, EXPANDING_NAME])
def test_ingestion_and_replacement_keep_identity_reads_working(  # noqa: PLR0913
    db_session: Session,
    provider: LocalKeyProvider,
    tmp_path: Path,
    connection_timing: str,
    withdraw: bool,
    initial_name: str,
) -> None:
    if connection_timing == "before":
        connect(db_session, provider)
    withdrawal: CanonicalWithdrawal | None = None
    mapping = MappingConfig(
        field_mappings={
            "counterparty": EntityMapping(
                source_table="Customers",
                fields={
                    "source_reference": "CustomerId",
                    "name": "CustomerName",
                    "counterparty_type": "Type",
                },
            ),
            "position": EntityMapping(
                source_table="Loans",
                fields={
                    "source_reference": "Reference",
                    "position_type": "Type",
                    "currency": "Currency",
                    "balance": "Balance",
                    "counterparty_reference": "CustomerId",
                },
            ),
        }
    )
    ingestion.create_mapping_config(
        db_session,
        CTX,
        BANK_A,
        MappingConfigCreate(
            source_system="EXCEL_CSV",
            name="Synthetic mapping",
            config=mapping,
            activate=True,
            reason="Synthetic encryption integration",
        ),
    )
    storage = InMemoryStorageClient()
    for name in (initial_name, "Synthetic replacement borrower"):
        path = tmp_path / "synthetic.xlsx"
        write_counterparty_workbook(path, name)
        result = ingestion.start_ingestion(
            db_session,
            CTX,
            BANK_A,
            IngestionBatchCreate(
                source_system="EXCEL_CSV",
                as_of_date=AS_OF,
                location=str(path),
                mapping_config_id=None,
                reason="Synthetic encrypted ingestion",
            ),
            storage,
        )
        assert result.batch.status in {"accepted", "accepted_with_warnings"}
        db_session.expunge_all()
        matches = list(
            db_session.scalars(
                select(CanonicalCounterparty)
                .options(undefer_group("bank_identity"))
                .where(
                    CanonicalCounterparty.organization_id == ORG_1,
                    CanonicalCounterparty.bank_id == BANK_A,
                    CanonicalCounterparty.superseded_by.is_(None),
                    CanonicalCounterparty.withdrawn_at.is_(None),
                )
            )
        )
        assert len(matches) == 1 and matches[0].name == name
        assert matches[0].source_reference == IDENTIFIER
        raw = cast(
            object,
            db_session.execute(
                text(
                    "SELECT name FROM canonical_counterparties "
                    "WHERE superseded_by IS NULL AND withdrawn_at IS NULL"
                )
            ).scalar_one(),
        )
        if connection_timing == "before" or (
            connection_timing == "after_withdrawal" and name != initial_name
        ):
            assert name not in str(raw) and str(raw).startswith("aqs-bank-field-v1:")
        else:
            assert raw == name
        if withdraw and name == initial_name:
            withdrawal = withdraw_counterparties(db_session)
            if connection_timing == "after_withdrawal":
                stored = raw_fields(db_session)
                connect(db_session, provider)
                assert raw_fields(db_session) == stored
    assert (
        db_session.execute(
            text("SELECT balance FROM canonical_position_snapshots WHERE superseded_by IS NULL")
        ).scalar_one()
        == 123
    )
    if withdrawal is not None:
        bank = db_session.get(Bank, BANK_A)
        assert bank is not None
        before_reversal = raw_fields(db_session)
        with pytest.raises(WithdrawalError, match="same natural key") as refused:
            canonical_withdrawal.reverse_withdrawal(
                db_session,
                CTX,
                bank,
                withdrawal.id,
                reversed_by="Synthetic checker",
                reason="Restore",
            )
        assert refused.value.status_code == 409
        db_session.flush()
        assert withdrawal.status == "applied"
        assert raw_fields(db_session) == before_reversal
        withdraw_counterparties(db_session)
        restored = canonical_withdrawal.reverse_withdrawal(
            db_session, CTX, bank, withdrawal.id, reversed_by="Synthetic checker", reason="Restore"
        )
        db_session.commit()
        assert restored.status == "reversed" and restored.rows_restored == 1
        matches = list(
            db_session.scalars(
                select(CanonicalCounterparty).where(
                    CanonicalCounterparty.organization_id == ORG_1,
                    CanonicalCounterparty.bank_id == BANK_A,
                    CanonicalCounterparty.source_system == "EXCEL_CSV",
                    CanonicalCounterparty.as_of_date == AS_OF,
                    CanonicalCounterparty.withdrawn_at.is_(None),
                    CanonicalCounterparty.superseded_by.is_(None),
                )
            )
        )
        assert len(matches) == 1 and matches[0].name == initial_name
        assert matches[0].source_reference == IDENTIFIER


@pytest.mark.parametrize("connection_timing", ["none", "before", "after_withdrawal"])
@pytest.mark.parametrize("axis", ["none", "reference", "source", "date", "bank"])
def test_counterparty_reversal_preserves_natural_key_scope(
    db_session: Session,
    provider: LocalKeyProvider,
    connection_timing: str,
    axis: str,
) -> None:
    if connection_timing == "before":
        connect(db_session, provider)
        connect(db_session, provider, BANK_B)
    original = counterparty(db_session, include_package=False)
    original_id = original.id
    withdrawal = withdraw_counterparties(db_session)
    if connection_timing == "after_withdrawal":
        connect(db_session, provider)
        connect(db_session, provider, BANK_B)
    replacement = None
    if axis != "none":
        replacement = counterparty(
            db_session,
            bank_id=BANK_B if axis == "bank" else BANK_A,
            reference="SYNTHETIC-OTHER" if axis == "reference" else IDENTIFIER,
            source_system="API_PUSH" if axis == "source" else "EXCEL_CSV",
            as_of=date(2026, 7, 31) if axis == "date" else AS_OF,
            include_package=False,
        )
        db_session.commit()
    bank = db_session.get(Bank, BANK_A)
    assert bank is not None
    restored = canonical_withdrawal.reverse_withdrawal(
        db_session, CTX, bank, withdrawal.id, reversed_by="Synthetic checker", reason="Restore"
    )
    db_session.commit()
    assert restored.status == "reversed" and restored.rows_restored == 1
    db_session.expire_all()
    loaded = db_session.get(CanonicalCounterparty, original_id)
    assert loaded is not None and loaded.withdrawn_at is None and loaded.name == NAME
    if replacement is not None:
        assert replacement.withdrawn_at is None and replacement.superseded_by is None


def test_connected_bank_bi_keeps_numeric_facts_without_identity_copies(
    db_session: Session,
    provider: LocalKeyProvider,
) -> None:

    connect(db_session, provider)
    party = counterparty(db_session, include_package=False)
    common = dict(
        organization_id=ORG_1,
        bank_id=BANK_A,
        as_of_date=AS_OF,
        source_system="EXCEL_CSV",
        ingestion_batch_id=party.ingestion_batch_id,
        lineage_id=party.lineage_id,
        validation_status="accepted",
    )
    position = CanonicalPosition(
        **common,
        source_reference="synthetic-deposit",
        position_type="DEPOSIT",
        currency="GHS",
    )
    db_session.add(position)
    db_session.flush()
    db_session.add(
        CanonicalPositionSnapshot(
            **common,
            source_reference="synthetic-snapshot",
            position_id=position.id,
            counterparty_id=party.id,
            balance=Decimal("123"),
        )
    )
    db_session.commit()
    provider.set_status(KEY_A, KeyStatus.DISABLED)
    db_session.expunge_all()
    outcome = materialize_bi_plane(db_session, organization_id=ORG_1, bank_id=BANK_A, as_of=AS_OF)
    assert outcome is not None and outcome.status == "succeeded"
    db_session.expunge_all()
    identity = db_session.execute(
        text("SELECT name, source_reference, group_reference FROM bi_dim_counterparty")
    ).one()
    assert NAME not in str(identity) and IDENTIFIER not in str(identity)
    assert identity[2] is None
    assert db_session.scalar(select(BiFactPositionDaily.counterparty_group)) is None
    assert db_session.scalar(select(BiFactPositionEom.counterparty_group)) is None
    q = BiQuery(measures=["deposits.balance_rc"], time=BiTime(as_of=AS_OF))
    compiled = compile_query(db_session, catalogue(), q, organization_id=ORG_1, bank_id=BANK_A)
    result = execute(db_session, compiled, timeout_ms=1000, row_cap=10)
    assert result.rows == [(Decimal("123"),)]
    for restricted in (
        q.model_copy(update={"dimensions": ["counterparty.name"]}),
        q.model_copy(update={"dimensions": ["counterparty"]}),
        q.model_copy(
            update={"filters": [BiFilter(member="counterparty.name", op="eq", values=[NAME])]}
        ),
    ):
        with pytest.raises(BiQueryError, match="bank-held encryption") as refused:
            compile_query(
                db_session, catalogue(), restricted, organization_id=ORG_1, bank_id=BANK_A
            )
        assert refused.value.code == "bank_key_identity_bi_unavailable"


def saved_version(db: Session, spec: dict[str, object]) -> BiDashboardVersion:
    dashboard = BiDashboard(
        organization_id=ORG_1,
        bank_id=BANK_A,
        owner_user_id=USER_1,
        title="Synthetic",
        description="",
        visibility="private",
        badge="personal",
        current_version=1,
    )
    db.add(dashboard)
    db.flush()
    version = BiDashboardVersion(
        organization_id=ORG_1,
        bank_id=BANK_A,
        dashboard_id=dashboard.id,
        version=1,
        title="Synthetic",
        description="",
        spec=spec,
        spec_digest="a" * 64,
        change_note="",
        created_by_user_id=USER_1,
        created_at=utc_now(),
    )
    return version


def test_saved_identity_filter_is_encrypted_and_explicitly_refused(
    db_session: Session,
    provider: LocalKeyProvider,
) -> None:

    connect(db_session, provider)
    saved_query = BiPackQuery(
        measures=["deposits.balance_rc"],
        window="as_of",
        filters=[BiFilter(member="counterparty.name", op="eq", values=[NAME])],
    )
    # Exercise the storage boundary; the product save gate below rejects
    # identity queries before accepting a user-authored canvas.
    version = saved_version(db_session, {"query": saved_query.model_dump(mode="json")})
    db_session.add(version)
    subscription = BiSubscription(
        organization_id=ORG_1,
        bank_id=BANK_A,
        name="Synthetic",
        owner_user_id=USER_1,
        query=saved_query.for_period(AS_OF).model_dump(mode="json"),
        artifact_format="csv",
        cadence="daily",
        hour=7,
        minute=30,
        recipient_user_ids=[],
    )
    db_session.add(subscription)
    db_session.commit()
    version_id = version.id
    assert version_id.version == 4
    raw = cast(
        object, db_session.execute(text("SELECT spec FROM bi_dashboard_versions")).scalar_one()
    )
    assert NAME not in str(raw) and "aqs-bank-field-v1" in str(raw)
    raw_query = cast(
        object, db_session.execute(text("SELECT query FROM bi_subscriptions")).scalar_one()
    )
    assert NAME not in str(raw_query) and "aqs-bank-field-v1" in str(raw_query)
    subscription_id = subscription.id
    db_session.expunge_all()
    restored_subscription = db_session.get(BiSubscription, subscription_id)
    assert restored_subscription is not None
    assert NAME in json.dumps(restored_subscription.query, ensure_ascii=False)
    restored = db_session.get(BiDashboardVersion, version_id)
    assert restored is not None
    assert NAME in json.dumps(restored.spec, ensure_ascii=False)
    spec = BiDashboardSpec.model_validate(
        {
            "widgets": [
                {
                    "id": "identity",
                    "kind": "table",
                    "title": "Synthetic",
                    "query": saved_query.model_dump(mode="json"),
                }
            ],
            "layout": [{"i": "identity", "x": 0, "y": 0, "w": 6, "h": 4}],
        }
    )
    bank = db_session.get(Bank, BANK_A)
    assert bank is not None
    with pytest.raises(content.MembersDenied) as refused:
        content.authorize_canvas(db_session, CTX, bank, spec, surface="test")
    assert refused.value.reason == IDENTITY_REASON
    authority = content.ViewerAuthority(
        db=db_session,
        ctx=CTX,
        bank=bank,
        cat=catalogue(),
        surface="test",
        _verdicts={("credit", "restricted"): True},
    )
    # A cached grant for another member of the same pair cannot allow identity.
    query = BiQuery(
        measures=["deposits.balance_rc"], time=BiTime(as_of=AS_OF), filters=saved_query.filters
    )
    assert authority.refused_members(query) == ("counterparty.name",)


def test_register_digest_tracks_stored_values(
    db_session: Session,
    provider: LocalKeyProvider,
) -> None:
    connect(db_session, provider)
    party = RelatedParty(
        organization_id=ORG_1,
        bank_id=BANK_A,
        party_type="individual",
        full_name=NAME,
        contact={"national_id": IDENTIFIER},
    )
    db_session.add(party)
    db_session.commit()
    party_id = party.id
    # Compare persisted projections on both sides. SQLite drops timezone info
    # on reload; a freshly-created in-memory timestamp is not that baseline.
    db_session.expunge_all()
    before = register_state_digest(register_state.project_register_state(db_session, CTX, BANK_A))
    db_session.expunge_all()
    after = register_state_digest(register_state.project_register_state(db_session, CTX, BANK_A))
    assert after == before
    loaded = db_session.get(RelatedParty, party_id)
    assert loaded is not None
    loaded.full_name = "Synthetic changed reportable name"
    db_session.commit()
    assert (
        register_state_digest(register_state.project_register_state(db_session, CTX, BANK_A))
        != before
    )


@pytest.mark.parametrize(
    "filters", [None, [], [{"member": "counterparty.name", "op": "eq", "values": [NAME]}]]
)
def test_saved_alert_list_filters_encrypt_and_round_trip(
    db_session: Session,
    provider: LocalKeyProvider,
    filters: list[dict[str, object]] | None,
) -> None:

    connect(db_session, provider)
    alert = BiAlert(
        organization_id=ORG_1,
        bank_id=BANK_A,
        owner_user_id=USER_1,
        name="Synthetic saved alert",
        measure_id="loans.balance_rc",
        filters=filters,
        direction="above",
        threshold_basis="stated",
        threshold=Decimal(100),
        notify_user_ids=[],
    )
    db_session.add(alert)
    db_session.commit()
    alert_id = alert.id
    db_session.expunge_all()
    raw = cast(object, db_session.execute(text("SELECT filters FROM bi_alerts")).scalar_one())
    assert ("aqs-bank-field-v1" in str(raw)) == bool(filters)
    assert NAME not in str(raw)
    loaded = db_session.get(BiAlert, alert_id)
    assert loaded is not None and loaded.filters == (filters or [])
    loaded.filters = [{"member": "counterparty.id", "op": "eq", "values": ["synthetic-id"]}]
    db_session.commit()
    db_session.expunge_all()
    updated = db_session.get(BiAlert, alert_id)
    assert updated is not None
    assert updated.filters == [
        {"member": "counterparty.id", "op": "eq", "values": ["synthetic-id"]}
    ]


def test_key_connection_encrypts_new_writes_without_rewriting_existing_records(
    db_session: Session,
    provider: LocalKeyProvider,
) -> None:
    previous = counterparty(db_session)
    previous_id = previous.id
    version = saved_version(db_session, {"borrower": NAME})
    db_session.add(version)
    db_session.commit()
    stored = raw_fields(db_session)
    connect(db_session, provider)
    assert raw_fields(db_session) == stored
    db_session.expunge_all()
    unchanged = db_session.get(CanonicalCounterparty, previous_id)
    assert unchanged is not None and unchanged.name == NAME
    current = counterparty(db_session, reference="SYNTHETIC-AFTER-CONNECTION")
    current_id = current.id
    db_session.commit()
    raw = cast(
        object,
        db_session.scalar(
            text("SELECT name FROM canonical_counterparties WHERE id=:id").bindparams(
                bindparam("id", type_=Uuid(as_uuid=True))
            ),
            {"id": current_id},
        ),
    )
    assert NAME not in str(raw) and "aqs-bank-field-v1" in str(raw)
    db_session.expunge_all()
    restored = db_session.get(CanonicalCounterparty, current_id)
    assert restored is not None and restored.name == NAME


@pytest.mark.parametrize("status", [KeyStatus.DISABLED, KeyStatus.UNAVAILABLE])
def test_numeric_book_and_filing_payloads_survive_key_loss(
    db_session: Session,
    provider: LocalKeyProvider,
    status: KeyStatus,
) -> None:
    connect(db_session, provider)
    party = counterparty(db_session, include_package=False)
    party.attributes = {"sovereign": True, "credit_conversion_factor": "0.5"}
    position = CanonicalPosition(
        organization_id=ORG_1,
        bank_id=BANK_A,
        as_of_date=AS_OF,
        source_system="EXCEL_CSV",
        source_reference="synthetic-credit",
        ingestion_batch_id=party.ingestion_batch_id,
        lineage_id=party.lineage_id,
        validation_status="accepted",
        position_type="LOAN",
        currency="GHS",
    )
    db_session.add(position)
    db_session.flush()
    db_session.add(
        CanonicalPositionSnapshot(
            organization_id=ORG_1,
            bank_id=BANK_A,
            as_of_date=AS_OF,
            source_system="EXCEL_CSV",
            source_reference="synthetic-credit-snapshot",
            ingestion_batch_id=party.ingestion_batch_id,
            lineage_id=party.lineage_id,
            validation_status="accepted",
            position_id=position.id,
            counterparty_id=party.id,
            balance=Decimal("123"),
        )
    )
    package = RegulatoryPackage(
        organization_id=ORG_1,
        bank_id=BANK_A,
        return_family="bsd",
        return_code="BSD1",
        reporting_date=AS_OF,
        frequency="monthly",
        status="generated",
        version=1,
        snapshot={"balance": "123"},
        validation_report={"valid": True},
        notes="Synthetic numeric filing",
        generated_by=USER_1,
    )
    db_session.add(package)
    db_session.flush()
    package_id, party_id = package.id, party.id
    db_session.commit()
    with Session(db_session.connection(), join_transaction_mode="create_savepoint") as fresh:
        bank = fresh.get(Bank, BANK_A)
        assert bank is not None
        before = position_book.credit_source_basis(
            position_book.load_position_records(fresh, CTX, bank, AS_OF)
        )
        expected_identity = fact_derivation.evaluate_balance_identity(fresh, CTX, bank, AS_OF)
        expected_exposures = credit_concentration.load_credit_exposures(fresh, CTX, bank, AS_OF)
        assert expected_identity is not None
        assert len(expected_exposures) == 1 and expected_exposures[0].ead == Decimal("123")
    provider.set_status(KEY_A, status)
    with Session(db_session.connection(), join_transaction_mode="create_savepoint") as fresh:
        bank = fresh.get(Bank, BANK_A)
        assert bank is not None
        records = position_book.load_position_records(fresh, CTX, bank, AS_OF)
        assert position_book.credit_source_basis(records) == before
        assert (
            fact_derivation.evaluate_balance_identity(fresh, CTX, bank, AS_OF) == expected_identity
        )
        assert (
            credit_concentration.load_credit_exposures(fresh, CTX, bank, AS_OF)
            == expected_exposures
        )
        loaded = fresh.get(RegulatoryPackage, package_id)
        assert loaded is not None and loaded.snapshot == {"balance": "123"}
        assert loaded.validation_report == {"valid": True}
        assert loaded.notes == "Synthetic numeric filing"
        loaded_party = fresh.get(CanonicalCounterparty, party_id)
        assert loaded_party is not None
        assert loaded_party.attributes["credit_conversion_factor"] == "0.5"
        with pytest.raises(KeyUnavailableError):
            _ = loaded_party.name


@pytest.mark.parametrize("connected", [False, True])
def test_mixed_bi_filters_protect_only_identity_values(
    db_session: Session,
    provider: LocalKeyProvider,
    connected: bool,
) -> None:
    if connected:
        connect(db_session, provider)
    numeric_filter = {"member": "position.currency", "op": "eq", "values": ["GHS"]}
    identity_filter = {"member": "counterparty.name", "op": "eq", "values": [NAME]}
    spec: dict[str, object] = {
        "widgets": [{"query": {"filters": [numeric_filter, identity_filter]}}]
    }
    version = saved_version(db_session, spec)
    db_session.add(version)
    db_session.flush()
    version_id = version.id
    db_session.commit()
    raw_value = cast(object, db_session.scalar(text("SELECT spec FROM bi_dashboard_versions")))
    raw = raw_value if isinstance(raw_value, str) else json.dumps(raw_value)
    stored = cast(dict[str, object], json.loads(raw))
    assert "GHS" in raw
    assert (NAME not in json.dumps(stored, ensure_ascii=False)) == connected
    assert ("aqs-bank-field-v1" in raw) == connected
    db_session.expunge_all()
    restored = db_session.get(BiDashboardVersion, version_id)
    assert restored is not None and restored.spec == spec
    db_session.commit()
    if connected:
        provider.set_status(KEY_A, KeyStatus.DISABLED)
        with Session(db_session.connection(), join_transaction_mode="create_savepoint") as fresh:
            locked = fresh.get(BiDashboardVersion, version_id)
            assert locked is not None and locked.spec == {LOCKED_JSON: "revoked"}
        provider.set_status(KEY_A, KeyStatus.ACTIVE)
        widgets = cast(list[dict[str, object]], stored["widgets"])
        query = cast(dict[str, object], widgets[0]["query"])
        filters = cast(list[object], query["filters"])
        filters.reverse()
        db_session.execute(
            text("UPDATE bi_dashboard_versions SET spec=:spec WHERE id=:id").bindparams(
                bindparam("id", type_=Uuid(as_uuid=True))
            ),
            {"spec": json.dumps(stored), "id": version_id},
        )
        db_session.commit()
        db_session.expunge_all()
        with pytest.raises(KeyUnavailableError):
            db_session.get(BiDashboardVersion, version_id)


@pytest.mark.parametrize("connected", [False, True])
def test_numeric_group_tokens_track_group_attributes_without_identity_reads(
    db_session: Session,
    provider: LocalKeyProvider,
    connected: bool,
) -> None:
    if connected:
        connect(db_session, provider)
    first = counterparty(db_session, include_package=False)
    second = counterparty(db_session, reference="synthetic-second", include_package=False)
    first.group_reference = second.group_reference = None
    first.attributes = {"parent": "Synthetic group A"}
    second.attributes = {"parent": "Synthetic group B"}
    db_session.commit()
    assert first.concentration_group_blind_index != second.concentration_group_blind_index
    second.attributes = {"parent": "Synthetic group A"}
    db_session.commit()
    expected = first.concentration_group_blind_index
    assert second.concentration_group_blind_index == expected
    second_id = second.id
    provider.set_status(KEY_A, KeyStatus.DISABLED)
    with Session(db_session.connection(), join_transaction_mode="create_savepoint") as fresh:
        loaded = fresh.get(CanonicalCounterparty, second_id)
        assert loaded is not None
        loaded.attributes = {"parent": "Synthetic group A", "sovereign": True}
        fresh.flush()
        assert loaded.attributes["sovereign"] is True
        assert loaded.concentration_group_blind_index == expected


@pytest.mark.parametrize("invalid_spec", [[], "invalid canvas"])
def test_numeric_saved_content_rejects_invalid_json_shape(
    db_session: Session,
    provider: LocalKeyProvider,
    invalid_spec: object,
) -> None:
    connect(db_session, provider)
    version = saved_version(db_session, {"widgets": []})
    db_session.add(version)
    db_session.flush()
    version_id = version.id
    db_session.commit()
    db_session.execute(
        text("UPDATE bi_dashboard_versions SET spec=:spec WHERE id=:id").bindparams(
            bindparam("id", type_=Uuid(as_uuid=True))
        ),
        {"spec": json.dumps(invalid_spec), "id": version_id},
    )
    db_session.commit()
    provider.set_status(KEY_A, KeyStatus.DISABLED)
    with (
        Session(db_session.connection(), join_transaction_mode="create_savepoint") as fresh,
        pytest.raises(KeyIntegrityError),
    ):
        fresh.get(BiDashboardVersion, version_id)


@pytest.mark.parametrize("status", [KeyStatus.DISABLED, KeyStatus.UNAVAILABLE])
def test_presentation_locks_and_recovers_without_changing_numeric_fields(
    db_session: Session,
    provider: LocalKeyProvider,
    status: KeyStatus,
) -> None:
    connect(db_session, provider)
    row = counterparty(db_session)
    row_id = row.id
    db_session.commit()
    provider.set_status(KEY_A, status)
    db_session.expunge_all()
    view = PresentationState()
    db_session.info[SESSION_STATE] = view
    locked = db_session.get(CanonicalCounterparty, row_id)
    assert locked is not None
    assert locked.name == LOCKED_TEXT
    assert locked.external_identifiers == {
        LOCKED_JSON: "revoked" if status == KeyStatus.DISABLED else "unavailable"
    }
    assert locked.as_of_date == AS_OF and locked.validation_status == "accepted"
    assert view.banks[BANK_A] == ("revoked" if status == KeyStatus.DISABLED else "unavailable")
    locked.name = "Synthetic edited name"
    with pytest.raises(KeyUnavailableError) as refused:
        db_session.flush()
    assert refused.value.revoked == (status == KeyStatus.DISABLED)
    db_session.rollback()
    provider.set_status(KEY_A, KeyStatus.ACTIVE)
    db_session.expunge_all()
    db_session.info[SESSION_STATE] = PresentationState()
    recovered = db_session.get(CanonicalCounterparty, row_id)
    assert recovered is not None and recovered.name == NAME
    assert recovered.attributes == {"contact_name": NAME}
    assert cast(PresentationState, db_session.info[SESSION_STATE]).banks[BANK_A] == "available"
    db_session.info.pop(SESSION_STATE)


@pytest.mark.parametrize("connected", [False, True])
@pytest.mark.parametrize("missing_presentation", [False, True])
def test_support_presentation_hides_fields_without_unwrapping(
    db_session: Session,
    provider: LocalKeyProvider,
    monkeypatch: pytest.MonkeyPatch,
    missing_presentation: bool,
    connected: bool,
) -> None:
    if connected:
        connect(db_session, provider)
    row_id = counterparty(db_session).id
    db_session.commit()
    db_session.expunge_all()

    def refuse_unwrap(*_args: object, **_kwargs: object) -> bytes:
        pytest.fail("Support presentation must not unwrap bank data")

    monkeypatch.setattr(provider, "unwrap", refuse_unwrap)
    db_session.info[SESSION_SUPPORT] = True
    if not missing_presentation:
        db_session.info[SESSION_STATE] = PresentationState(support=True)
    row = db_session.get(CanonicalCounterparty, row_id)
    assert row is not None and row.name == LOCKED_TEXT
    if not missing_presentation:
        assert cast(PresentationState, db_session.info[SESSION_STATE]).banks == {BANK_A: "support"}
    db_session.info.pop(SESSION_STATE, None)
    db_session.info.pop(SESSION_SUPPORT)


@pytest.mark.parametrize("status", [KeyStatus.DISABLED, KeyStatus.UNAVAILABLE])
@pytest.mark.parametrize("legacy", [False, True])
def test_profile_http_locked_read_write_and_recovery(
    db_client: TestClient,
    db_session: Session,
    provider: LocalKeyProvider,
    status: KeyStatus,
    legacy: bool,
) -> None:
    if not legacy:
        connect(db_session, provider)
    party = institution_profile.create_related_party(
        db_session,
        CTX,
        BANK_A,
        RelatedPartyCreate(
            reason="Synthetic protected profile",
            party_type="individual",
            full_name=NAME,
            contact={"email": "synthetic@example.test"},
            roles=[RelatedPartyRoleInput(role="director")],
        ),
    )
    db_session.commit()
    if legacy:
        connect(db_session, provider)
    provider.set_status(KEY_A, status)
    db_session.expunge_all()
    response = db_client.get(f"/api/v1/banks/{BANK_A}/institution-profile", headers=headers())
    assert response.status_code == 200, response.text
    assert response.json()["related_parties"][0]["full_name"] == LOCKED_TEXT
    reason = "revoked" if status == KeyStatus.DISABLED else "unavailable"
    assert json.loads(response.headers["X-Aequoros-Protected-Data"]) == {BANK_A: reason}
    assert "no-store" in response.headers["cache-control"]
    payload: dict[str, object] = {
        "reason": "Synthetic retained edit",
        "party_type": "individual",
        "full_name": "Synthetic draft stays in browser",
        "contact": {},
        "roles": [{"role": "director"}],
    }
    refused = db_client.put(
        f"/api/v1/banks/{BANK_A}/related-parties/{party.id}",
        headers=headers(),
        json=payload,
    )
    assert refused.status_code == 503, refused.text
    assert (
        refused.json()["error"]["message"]
        == KeyUnavailableError(
            "private provider detail",
            revoked=status == KeyStatus.DISABLED,
        ).public_message
    )
    provider.set_status(KEY_A, KeyStatus.ACTIVE)
    restored = db_client.get(f"/api/v1/banks/{BANK_A}/institution-profile", headers=headers())
    assert restored.status_code == 200, restored.text
    assert restored.json()["related_parties"][0]["full_name"] == NAME
    assert json.loads(restored.headers["X-Aequoros-Protected-Data"]) == {BANK_A: "available"}


@pytest.mark.parametrize("status", [KeyStatus.DISABLED, KeyStatus.UNAVAILABLE])
def test_empty_profile_observes_key_recovery_after_a_refused_party_create(
    db_client: TestClient,
    db_session: Session,
    provider: LocalKeyProvider,
    status: KeyStatus,
) -> None:
    connect(db_session, provider)
    url = f"/api/v1/banks/{BANK_A}/institution-profile"
    provider.set_status(KEY_A, status)
    payload: dict[str, object] = {
        "reason": "Synthetic retained draft",
        "party_type": "individual",
        "full_name": NAME,
        "contact": {},
        "roles": [],
    }
    refused = db_client.post(
        f"/api/v1/banks/{BANK_A}/related-parties", headers=headers(), json=payload
    )
    assert refused.status_code == 503, refused.text
    reason = "revoked" if status == KeyStatus.DISABLED else "unavailable"
    for _ in range(2):
        empty = db_client.get(url, headers=headers())
        assert empty.status_code == 200, empty.text
        assert empty.json()["related_parties"] == []
        assert json.loads(empty.headers["X-Aequoros-Protected-Data"]) == {BANK_A: reason}
    provider.set_status(KEY_A, KeyStatus.ACTIVE)
    restored = db_client.get(url, headers=headers())
    assert restored.status_code == 200, restored.text
    assert restored.json()["related_parties"] == []
    assert json.loads(restored.headers["X-Aequoros-Protected-Data"]) == {BANK_A: "available"}
    saved = db_client.post(
        f"/api/v1/banks/{BANK_A}/related-parties", headers=headers(), json=payload
    )
    assert saved.status_code == 201, saved.text
    assert saved.json()["full_name"] == NAME


@pytest.mark.parametrize("missing_current", [False, True])
@pytest.mark.parametrize("status", [KeyStatus.DISABLED, KeyStatus.UNAVAILABLE])
def test_rotated_profile_keeps_failure_until_both_write_keys_recover(
    db_client: TestClient,
    db_session: Session,
    provider: LocalKeyProvider,
    missing_current: bool,
    status: KeyStatus,
) -> None:
    connect(db_session, provider)
    payload = RelatedPartyCreate(
        reason="Synthetic rotated recovery",
        party_type="individual",
        full_name="Synthetic retained-key party",
        contact={"email": "retained@example.test"},
        roles=[RelatedPartyRoleInput(role="director", sitting_allowance=Decimal(45))],
    )
    retained = institution_profile.create_related_party(db_session, CTX, BANK_A, payload)
    row = registry.scoped_key(db_session, bank_id=BANK_A, organization_id=ORG_1)
    registry.rotate(db_session, row, NEXT, lambda _name: provider)
    assert row.current_version_id != row.index_version_id
    db_session.commit()
    current = institution_profile.create_related_party(
        db_session,
        CTX,
        BANK_A,
        payload.model_copy(update={"full_name": "Synthetic current-key party"}),
    )
    stored_query = text("SELECT full_name FROM related_parties WHERE bank_id=:bank ORDER BY id")
    stored = TypeAdapter(list[str]).validate_python(
        cast(object, db_session.scalars(stored_query, {"bank": BANK_A}).all())
    )
    db_session.commit()
    missing = NEXT if missing_current else KEY_A
    provider.set_status(missing, status)
    url = f"/api/v1/banks/{BANK_A}/institution-profile"
    write_url = f"/api/v1/banks/{BANK_A}/related-parties"
    refused = db_client.post(write_url, headers=headers(), json=payload.model_dump(mode="json"))
    assert refused.status_code == 503, refused.text
    reason = "revoked" if status == KeyStatus.DISABLED else "unavailable"
    assert refused.json()["error"]["code"] == f"bank_key_{reason}"
    for _ in range(2):
        response = db_client.get(url, headers=headers())
        assert response.status_code == 200, response.text
        assert json.loads(response.headers["X-Aequoros-Protected-Data"]) == {BANK_A: reason}
        parties = {
            party.id: party
            for party in InstitutionProfileFullRead.model_validate_json(
                response.content
            ).related_parties
        }
        visible, locked = (retained, current) if missing_current else (current, retained)
        assert parties[visible.id].full_name == visible.full_name
        assert parties[visible.id].contact == {"email": "retained@example.test"}
        assert parties[locked.id].full_name == LOCKED_TEXT
        assert all(party.roles[0].sitting_allowance == Decimal(45) for party in parties.values())
    assert (
        TypeAdapter(list[str]).validate_python(
            cast(object, db_session.scalars(stored_query, {"bank": BANK_A}).all())
        )
        == stored
    )
    db_session.commit()
    provider.set_status(missing, KeyStatus.ACTIVE)
    recovered = db_client.get(url, headers=headers())
    assert recovered.status_code == 200, recovered.text
    assert json.loads(recovered.headers["X-Aequoros-Protected-Data"]) == {BANK_A: "available"}
    restored = InstitutionProfileFullRead.model_validate_json(recovered.content)
    assert {party.full_name for party in restored.related_parties} == {
        retained.full_name,
        current.full_name,
    }
    saved = db_client.post(write_url, headers=headers(), json=payload.model_dump(mode="json"))
    assert saved.status_code == 201, saved.text
    assert saved.json()["full_name"] == payload.full_name


@pytest.mark.parametrize("rotated", [False, True])
def test_profile_recovery_unwraps_each_required_version_once(
    db_client: TestClient,
    db_session: Session,
    provider: LocalKeyProvider,
    monkeypatch: pytest.MonkeyPatch,
    rotated: bool,
) -> None:
    connect(db_session, provider)
    for index in range(3):
        institution_profile.create_related_party(
            db_session,
            CTX,
            BANK_A,
            RelatedPartyCreate(
                reason="Synthetic bounded key observation",
                party_type="individual",
                full_name=f"Synthetic historical party {index}",
                contact={"email": "synthetic@example.test"},
                roles=[RelatedPartyRoleInput(role="director", other_responsibilities="Review")],
            ),
        )
    if rotated:
        row = registry.scoped_key(db_session, bank_id=BANK_A, organization_id=ORG_1)
        registry.rotate(db_session, row, NEXT, lambda _name: provider)
        db_session.commit()
    unwrap = provider.unwrap
    calls: list[KeyReference] = []

    def observed(key: KeyReference, wrapped: bytes, context: dict[str, str]) -> bytes:
        calls.append(key)
        return unwrap(key, wrapped, context)

    monkeypatch.setattr(provider, "unwrap", observed)
    response = db_client.get(f"/api/v1/banks/{BANK_A}/institution-profile", headers=headers())
    assert response.status_code == 200, response.text
    assert json.loads(response.headers["X-Aequoros-Protected-Data"]) == {BANK_A: "available"}
    assert (
        len(InstitutionProfileFullRead.model_validate_json(response.content).related_parties) == 3
    )
    assert calls.count(KEY_A) == 1 and calls.count(NEXT) == int(rotated)
    assert len(calls) == 1 + int(rotated)


@pytest.mark.parametrize("populated", [False, True])
@pytest.mark.parametrize(
    "scenario",
    [
        (KeyStatus.UNAVAILABLE, KeyStatus.DISABLED, NEXT),
        (KeyStatus.DISABLED, KeyStatus.UNAVAILABLE, KEY_A),
    ],
)
def test_rotated_profiles_combine_prerequisite_failure_severity(
    db_client: TestClient,
    db_session: Session,
    provider: LocalKeyProvider,
    populated: bool,
    scenario: tuple[KeyStatus, KeyStatus, KeyReference],
) -> None:
    connect(db_session, provider)
    if populated:
        institution_profile.create_related_party(
            db_session,
            CTX,
            BANK_A,
            RelatedPartyCreate(
                reason="Synthetic historical field failure",
                party_type="individual",
                full_name=NAME,
                contact={"email": "synthetic@example.test"},
                roles=[],
            ),
        )
    row = registry.scoped_key(db_session, bank_id=BANK_A, organization_id=ORG_1)
    registry.rotate(db_session, row, NEXT, lambda _name: provider)
    db_session.commit()
    index_status, current_status, first_restored = scenario
    provider.set_status(NEXT, current_status)
    provider.set_status(KEY_A, index_status)
    url = f"/api/v1/banks/{BANK_A}/institution-profile"
    for expected in ("revoked", "unavailable", "available"):
        response = db_client.get(url, headers=headers())
        assert response.status_code == 200, response.text
        parties = InstitutionProfileFullRead.model_validate_json(response.content).related_parties
        assert len(parties) == int(populated)
        if populated:
            assert parties[0].full_name == (
                NAME
                if expected == "available" or first_restored == KEY_A and expected == "unavailable"
                else LOCKED_TEXT
            )
        assert json.loads(response.headers["X-Aequoros-Protected-Data"]) == {BANK_A: expected}
        provider.set_status(first_restored if expected == "revoked" else KEY_A, KeyStatus.ACTIVE)
        if expected == "unavailable":
            provider.set_status(NEXT, KeyStatus.ACTIVE)


@pytest.mark.parametrize("observe_first", [False, True])
@pytest.mark.parametrize(
    "scenario",
    [
        (KeyStatus.ACTIVE, KeyStatus.UNAVAILABLE, KeyStatus.DISABLED, "revoked"),
        (KeyStatus.DISABLED, KeyStatus.ACTIVE, KeyStatus.UNAVAILABLE, "revoked"),
        (KeyStatus.UNAVAILABLE, KeyStatus.DISABLED, KeyStatus.ACTIVE, "revoked"),
        (KeyStatus.ACTIVE, KeyStatus.UNAVAILABLE, KeyStatus.ACTIVE, "unavailable"),
        (KeyStatus.UNAVAILABLE, KeyStatus.ACTIVE, KeyStatus.ACTIVE, "unavailable"),
        (KeyStatus.ACTIVE, KeyStatus.ACTIVE, KeyStatus.UNAVAILABLE, "unavailable"),
        (KeyStatus.ACTIVE, KeyStatus.ACTIVE, KeyStatus.ACTIVE, "available"),
    ],
)
def test_prerequisite_observation_is_independent_of_historical_read_order(
    db_session: Session,
    provider: LocalKeyProvider,
    monkeypatch: pytest.MonkeyPatch,
    observe_first: bool,
    scenario: tuple[KeyStatus, KeyStatus, KeyStatus, str],
) -> None:
    connect(db_session, provider)
    row = registry.scoped_key(db_session, bank_id=BANK_A, organization_id=ORG_1)
    registry.rotate(db_session, row, NEXT, lambda _name: provider)
    db_session.commit()
    institution_profile.create_related_party(
        db_session,
        CTX,
        BANK_A,
        RelatedPartyCreate(
            reason="Synthetic independent historical key",
            party_type="individual",
            full_name=NAME,
            contact={"email": "synthetic@example.test"},
            roles=[RelatedPartyRoleInput(role="director", sitting_allowance=Decimal(45))],
        ),
    )
    current_key = KeyReference("aws_kms", "synthetic-current-v3", "us-east-1", "111122223333")
    provider.add_key(current_key)
    row = registry.scoped_key(db_session, bank_id=BANK_A, organization_id=ORG_1)
    registry.rotate(db_session, row, current_key, lambda _name: provider)
    db_session.commit()
    index_status, historical_status, current_status, expected = scenario
    for key, status in (
        (KEY_A, index_status),
        (NEXT, historical_status),
        (current_key, current_status),
    ):
        provider.set_status(key, status)
    calls: list[KeyReference] = []
    unwrap = provider.unwrap

    def observed(key: KeyReference, wrapped: bytes, context: dict[str, str]) -> bytes:
        calls.append(key)
        return unwrap(key, wrapped, context)

    monkeypatch.setattr(provider, "unwrap", observed)
    with Session(db_session.connection(), join_transaction_mode="create_savepoint") as fresh:
        presentation = PresentationState()
        fresh.info[SESSION_STATE] = presentation
        if observe_first:
            observe_bank_key(fresh, ORG_1, BANK_A)
        profile = institution_profile.get_full_profile(fresh, CTX, BANK_A)
        observe_bank_key(fresh, ORG_1, BANK_A)
        assert presentation.banks == {BANK_A: expected}
        assert profile.related_parties[0].full_name == (
            NAME if historical_status == KeyStatus.ACTIVE else LOCKED_TEXT
        )
        assert profile.related_parties[0].roles[0].sitting_allowance == Decimal(45)
        assert all(calls.count(key) == 1 for key in (KEY_A, NEXT, current_key))
        assert len(calls) == 3
    for key in (KEY_A, NEXT, current_key):
        provider.set_status(key, KeyStatus.ACTIVE)
    calls.clear()
    with Session(db_session.connection(), join_transaction_mode="create_savepoint") as fresh:
        presentation = PresentationState()
        fresh.info[SESSION_STATE] = presentation
        profile = institution_profile.get_full_profile(fresh, CTX, BANK_A)
        assert presentation.banks == {BANK_A: "available"}
        assert profile.related_parties[0].full_name == NAME
        assert all(calls.count(key) == 1 for key in (KEY_A, NEXT, current_key))
        assert len(calls) == 3


@pytest.mark.parametrize("connected", [False, True])
def test_inspector_profile_read_never_releases_protected_identity(
    db_client: TestClient,
    db_session: Session,
    provider: LocalKeyProvider,
    monkeypatch: pytest.MonkeyPatch,
    connected: bool,
) -> None:
    if connected:
        connect(db_session, provider)
    institution_profile.create_related_party(
        db_session,
        CTX,
        BANK_A,
        RelatedPartyCreate(
            reason="Synthetic staff disclosure proof",
            party_type="individual",
            full_name=NAME,
            contact={"email": "identity@example.test"},
            roles=[RelatedPartyRoleInput(role="director")],
        ),
    )
    db_session.commit()

    def refuse_unwrap(*_args: object, **_kwargs: object) -> bytes:
        pytest.fail("Staff reads must not unwrap identity fields")

    monkeypatch.setattr(provider, "unwrap", refuse_unwrap)
    token = impersonation._impersonation_token(ORG_1)  # pyright: ignore[reportPrivateUsage]
    response = db_client.get(
        f"/api/v1/banks/{BANK_A}/institution-profile", headers={"Authorization": f"Bearer {token}"}
    )
    assert response.status_code == 200, response.text
    assert response.json()["related_parties"][0]["full_name"] == LOCKED_TEXT
    assert NAME not in response.text and "identity@example.test" not in response.text
    assert json.loads(response.headers["X-Aequoros-Protected-Data"]) == {BANK_A: "support"}


def test_corrupted_protected_read_is_an_integrity_refusal_not_an_outage(
    db_client: TestClient,
    db_session: Session,
    provider: LocalKeyProvider,
) -> None:
    connect(db_session, provider)
    party = institution_profile.create_related_party(
        db_session,
        CTX,
        BANK_A,
        RelatedPartyCreate(
            reason="Synthetic integrity proof",
            party_type="individual",
            full_name=NAME,
            contact={},
            roles=[RelatedPartyRoleInput(role="director")],
        ),
    )
    db_session.execute(
        text("UPDATE related_parties SET full_name=:value WHERE id=:id").bindparams(
            bindparam("id", type_=Uuid(as_uuid=True))
        ),
        {"value": "aqs-bank-field-v1:invalid", "id": party.id},
    )
    db_session.commit()
    response = db_client.get(f"/api/v1/banks/{BANK_A}/institution-profile", headers=headers())
    assert response.status_code == 409, response.text
    assert response.json()["error"]["code"] == "bank_key_integrity"
    assert "Retrying" not in response.text


def group_snapshot(
    db: Session,
    party: CanonicalCounterparty,
    reference: str,
    as_of: date,
    position_type: str,
) -> CanonicalPositionSnapshot:
    common = dict(
        organization_id=ORG_1,
        bank_id=party.bank_id,
        as_of_date=as_of,
        source_system="EXCEL_CSV",
        ingestion_batch_id=party.ingestion_batch_id,
        lineage_id=party.lineage_id,
        validation_status="accepted",
    )
    position = CanonicalPosition(
        **common, source_reference=reference, position_type=position_type, currency="GHS"
    )
    db.add(position)
    db.flush()
    snapshot = CanonicalPositionSnapshot(
        **common,
        source_reference=reference,
        position_id=position.id,
        counterparty_id=party.id,
        balance=Decimal(100),
    )
    db.add(snapshot)
    db.flush()
    return snapshot


@pytest.mark.parametrize("connected", [False, True])
@pytest.mark.parametrize("length", [249, 255])
@pytest.mark.parametrize(
    "representation", ["name", "group_reference", "group", "parent", "attribute_reference"]
)
def test_long_group_tokens_preserve_numeric_grouping(
    db_session: Session,
    provider: LocalKeyProvider,
    connected: bool,
    length: int,
    representation: str,
) -> None:
    if connected:
        connect(db_session, provider)
    labels: list[str] = []
    for index in range(2):
        value = "S" * (length - 1) + str(index)
        party = counterparty(db_session, reference=f"long-party-{index}", name=value)
        party.group_reference = None
        party.attributes = {}
        if representation == "group_reference":
            party.group_reference = value
        elif representation != "name":
            key = "group_reference" if representation == "attribute_reference" else representation
            party.attributes = {key: value}
        labels.append(counterparty_group_key(party.name, party.group_reference, party.attributes))
        group_snapshot(db_session, party, f"long-loan-{index}", AS_OF, "LOAN")
    db_session.commit()
    db_session.expunge_all()
    tokens = list(db_session.scalars(select(CanonicalCounterparty.concentration_group_blind_index)))
    assert len(tokens) == 2 and all(token is not None and len(token) <= 255 for token in tokens)
    assert len(set(tokens)) == 2
    for label in labels:
        assert (
            blind_index(
                db_session,
                organization_id=ORG_1,
                bank_id=BANK_A,
                domain="counterparty.group",
                value=label,
            )
            in tokens
        )
    bank = db_session.get(Bank, BANK_A)
    assert bank is not None
    report = credit_concentration.concentration_read(db_session, CTX, bank, AS_OF)
    single_name = next(d for d in report.dimensions if d.dimension == "single_name")
    assert {bucket.key for bucket in single_name.buckets} == set(labels)
    assert single_name.hhi == Decimal(5000)


@pytest.mark.parametrize(
    "representation", ["name", "group_reference", "group", "parent", "attribute_reference"]
)
def test_optional_bank_named_limit_matches_the_stored_group(
    db_session: Session,
    provider: LocalKeyProvider,
    representation: str,
) -> None:
    _credit_identity_positions(db_session, representation)
    label = "cp:Synthetic borrower 0" if representation == "name" else "group:Synthetic group 0"
    db_session.add(
        ParamConcentrationLimit(
            organization_id=ORG_1,
            jurisdiction_code="GH",
            dimension="single_name",
            limit_kind="share_of_book_pct",
            value=Decimal(49),
            effective_from=AS_OF,
            approved_by="Synthetic board",
            approval_timestamp=utc_now(),
            bucket_key=label,
        )
    )
    db_session.commit()
    bank = db_session.get(Bank, BANK_A)
    assert bank is not None
    report = credit_concentration.concentration_read(db_session, CTX, bank, AS_OF)
    assert len(report.breaches) == 1 and report.breaches[0].key == label


@pytest.mark.parametrize("connected", [False, True])
@pytest.mark.parametrize("position_type", ["LOAN", "DEPOSIT"])
def test_report_labels_only_read_needed_snapshot_representatives(
    db_session: Session,
    provider: LocalKeyProvider,
    connected: bool,
    position_type: str,
) -> None:
    if connected:
        connect(db_session, provider)
    selected = counterparty(db_session, reference="selected", as_of=date(2020, 1, 1))
    selected.group_reference = None
    selected.attributes = {"parent": "Synthetic representative"}
    group_snapshot(db_session, selected, "selected-position", AS_OF, position_type)
    selected_id = selected.id
    excluded_ids: list[UUID] = []
    for index, reason in enumerate(
        (
            "old",
            "future",
            "unreferenced",
            "withdrawn",
            "superseded",
            "invalid",
            "wrong_type",
            "wrong_bank",
            "older_group",
            "duplicate_group",
        )
    ):
        party = counterparty(
            db_session,
            bank_id=BANK_B if reason == "wrong_bank" else BANK_A,
            reference=f"excluded-{index}",
            as_of=AS_OF,
        )
        if reason in {"older_group", "duplicate_group"}:
            party.group_reference = None
            party.attributes = {"parent": "Synthetic representative"}
        observed = (
            date(2020, 1, 1)
            if reason == "old"
            else date(2026, 7, 1)
            if reason == "future"
            else date(2026, 5, 30)
            if reason == "older_group"
            else AS_OF
        )
        if reason != "unreferenced" and not (reason == "future" and position_type == "DEPOSIT"):
            snapshot = group_snapshot(
                db_session,
                party,
                f"excluded-position-{index}",
                observed,
                ("DEPOSIT" if position_type == "LOAN" else "LOAN")
                if reason == "wrong_type"
                else position_type,
            )
            if reason == "withdrawn":
                snapshot.withdrawn_at = utc_now()
                snapshot.withdrawn_by_batch_id = party.ingestion_batch_id
                snapshot.withdrawal_reason = "Synthetic withdrawal"
            elif reason == "superseded":
                replacement = group_snapshot(
                    db_session, selected, "replacement-position", AS_OF, position_type
                )
                snapshot.superseded_by = replacement.id
            elif reason == "invalid":
                snapshot.validation_status = "error"
        excluded_ids.append(party.id)
    db_session.commit()
    db_session.execute(
        text(
            "UPDATE canonical_counterparties SET source_reference = :broken, "
            "external_identifiers = :json WHERE id = :id"
        ).bindparams(bindparam("id", type_=Uuid())),
        {
            "id": selected_id,
            "broken": "aqs-bank-field-v1:[]",
            "json": json.dumps({"aqs-bank-field-v1": []}),
        },
    )
    for record_id in excluded_ids:
        db_session.execute(
            text("UPDATE canonical_counterparties SET name = :broken WHERE id = :id").bindparams(
                bindparam("id", type_=Uuid())
            ),
            {"id": record_id, "broken": "aqs-bank-field-v1:[]"},
        )
    db_session.commit()
    db_session.expunge_all()
    if position_type == "LOAN":
        bank = db_session.get(Bank, BANK_A)
        assert bank is not None
        report = credit_concentration.concentration_read(db_session, CTX, bank, AS_OF)
        single_name = next(d for d in report.dimensions if d.dimension == "single_name")
        assert {bucket.key for bucket in single_name.buckets} == {"group:Synthetic representative"}
        assert single_name.bucket_count == 1
        assert report.total_book_ghs == Decimal(300)
    else:
        report = get_behavioral_liquidity_report(db_session, CTX, BANK_A)
        groups = [
            segment for segment in report.segments if segment.dimension == "concentration_group"
        ]
        assert {segment.segment for segment in groups} == {"Synthetic representative"}
        assert len(groups) == 1
        assert groups[0].current_balance_ghs == 300


@pytest.mark.parametrize("status", [KeyStatus.DISABLED, KeyStatus.UNAVAILABLE])
def test_locked_groups_remain_distinct_and_numeric_training_stays_available(
    db_session: Session, provider: LocalKeyProvider, status: KeyStatus
) -> None:
    connect(db_session, provider)
    parties = [counterparty(db_session), counterparty(db_session, reference="SYNTHETIC-SECOND")]
    parties[1].group_reference = "Synthetic separate group"
    for index, party in enumerate(parties):
        common = dict(
            organization_id=ORG_1,
            bank_id=BANK_A,
            source_system="EXCEL_CSV",
            ingestion_batch_id=party.ingestion_batch_id,
            lineage_id=party.lineage_id,
            validation_status="accepted",
        )
        position = CanonicalPosition(
            **common,
            as_of_date=AS_OF,
            source_reference=f"deposit-{index}",
            position_type="DEPOSIT",
            currency="GHS",
        )
        db_session.add(position)
        db_session.flush()
        for month, balance in ((5, 100 if index == 0 else 0), (6, 0 if index == 0 else 100)):
            db_session.add(
                CanonicalPositionSnapshot(
                    **common,
                    as_of_date=date(2026, month, 30),
                    source_reference=f"snapshot-{index}-{month}",
                    position_id=position.id,
                    counterparty_id=party.id,
                    balance=Decimal(balance),
                )
            )
    db_session.commit()
    db_session.expunge_all()
    provider.set_status(KEY_A, status)
    for estimate in (deposit_stability.estimate, nmd_duration.estimate):
        result = estimate(db_session, CTX, BANK_A, AS_OF, BehavioralTrainingConfig())
        assert result.products
    db_session.info[SESSION_STATE] = PresentationState()
    report = get_behavioral_liquidity_report(db_session, CTX, BANK_A)
    groups = [item for item in report.segments if item.dimension == "concentration_group"]
    assert len(groups) == 2
    assert {item.segment for item in groups} == {LOCKED_TEXT}
    assert sorted(item.current_balance_ghs for item in groups) == [0, 100]
    assert (
        next(item for item in groups if item.current_balance_ghs == 0).observed_monthly_runoff_pct
        == 100
    )
    db_session.info.pop(SESSION_STATE)


def _credit_identity_positions(db_session: Session, representation: str) -> None:
    for index in range(2):
        party = counterparty(
            db_session, reference=f"credit-party-{index}", name=f"Synthetic borrower {index}"
        )
        party.group_reference = None
        party.attributes = {}
        if representation == "group_reference":
            party.group_reference = f"Synthetic group {index}"
        elif representation != "name":
            key = "group_reference" if representation == "attribute_reference" else representation
            party.attributes = {key: f"Synthetic group {index}"}
        common = dict(
            organization_id=ORG_1,
            bank_id=BANK_A,
            as_of_date=AS_OF,
            source_system="EXCEL_CSV",
            ingestion_batch_id=party.ingestion_batch_id,
            lineage_id=party.lineage_id,
            validation_status="accepted",
        )
        position = CanonicalPosition(
            **common,
            source_reference=f"credit-position-{index}",
            position_type="LOAN",
            currency="GHS",
        )
        db_session.add(position)
        db_session.flush()
        db_session.add(
            CanonicalPositionSnapshot(
                **common,
                source_reference=f"credit-snapshot-{index}",
                position_id=position.id,
                counterparty_id=party.id,
                balance=Decimal(100),
            )
        )
    db_session.commit()
    db_session.expunge_all()


@pytest.mark.parametrize(
    "representation", ["group_reference", "attribute_reference", "group", "parent", "name"]
)
@pytest.mark.parametrize("status", [KeyStatus.DISABLED, KeyStatus.UNAVAILABLE])
def test_credit_concentration_identity_survives_locked_labels(
    db_session: Session,
    provider: LocalKeyProvider,
    representation: str,
    status: KeyStatus,
) -> None:
    connect(db_session, provider)
    _credit_identity_positions(db_session, representation)
    bank = db_session.get(Bank, BANK_A)
    assert bank is not None
    before = credit_concentration.load_credit_exposures(db_session, CTX, bank, AS_OF)
    shared_before = credit_exposure_book.load_exposure_rows(
        db_session,
        CTX,
        bank,
        AS_OF,
        credit_exposure_book.CONCENTRATION_POSITION_TYPES,
    )
    assert len({item.group_key for item in before}) == 2
    limit = ParamConcentrationLimit(
        organization_id=ORG_1,
        jurisdiction_code="GH",
        dimension="single_name",
        limit_kind="share_of_book_pct",
        value=Decimal(49),
        effective_from=AS_OF,
        approved_by="Synthetic board",
        approval_timestamp=utc_now(),
        bucket_key="cp:Synthetic borrower 0"
        if representation == "name"
        else "group:Synthetic group 0",
    )
    db_session.add(limit)
    db_session.flush()
    limited = credit_concentration.monitor(db_session, CTX, bank, AS_OF)
    assert len(limited.breaches) == 1 and limited.breaches[0].key == before[0].group_key
    db_session.commit()
    db_session.expunge_all()
    provider.set_status(KEY_A, status)
    after = credit_concentration.load_credit_exposures(db_session, CTX, bank, AS_OF)
    assert after == before
    monitor = credit_concentration.monitor(db_session, CTX, bank, AS_OF)
    dimension = monitor.dimension("single_name")
    assert dimension is not None and dimension.hhi == Decimal(5000)
    assert [item.share_of_book_pct for item in dimension.buckets] == [Decimal(50), Decimal(50)]
    db_session.info[SESSION_STATE] = PresentationState()
    shared_after = credit_exposure_book.load_exposure_rows(
        db_session,
        CTX,
        bank,
        AS_OF,
        credit_exposure_book.CONCENTRATION_POSITION_TYPES,
    )
    assert [item.group_key for item in shared_after] == [item.group_key for item in shared_before]
    report = credit_concentration.concentration_read(db_session, CTX, bank, AS_OF)
    displayed = next(item for item in report.dimensions if item.dimension == "single_name")
    assert len(displayed.buckets) == 2 and {item.key for item in displayed.buckets} == {LOCKED_TEXT}
    assert report.limit_count == 1
    assert len(report.locked_limits) == 1
    assert report.locked_limits[0].limit_status == "locked"
    assert not report.breaches
    assert all(item.limit_status == "locked" for item in displayed.buckets)
    provider.set_status(KEY_A, KeyStatus.ACTIVE)
    db_session.expunge_all()
    recovered = credit_concentration.concentration_read(db_session, CTX, bank, AS_OF)
    assert not recovered.locked_limits
    assert len(recovered.breaches) == 1
    db_session.info.pop(SESSION_STATE)


def test_positional_concentration_limit_matches_without_key(
    db_session: Session, provider: LocalKeyProvider
) -> None:
    connect(db_session, provider)
    party = counterparty(db_session)
    common = dict(
        organization_id=ORG_1,
        bank_id=BANK_A,
        as_of_date=AS_OF,
        source_system="EXCEL_CSV",
        ingestion_batch_id=party.ingestion_batch_id,
        lineage_id=party.lineage_id,
        validation_status="accepted",
    )
    position = CanonicalPosition(
        **common, source_reference="Synthetic position", position_type="LOAN", currency="GHS"
    )
    db_session.add(position)
    db_session.flush()
    db_session.add(
        CanonicalPositionSnapshot(
            **common,
            source_reference="Synthetic snapshot",
            position_id=position.id,
            balance=Decimal(100),
        )
    )
    db_session.add(
        ParamConcentrationLimit(
            organization_id=ORG_1,
            jurisdiction_code="GH",
            dimension="single_name",
            limit_kind="share_of_book_pct",
            value=Decimal(49),
            effective_from=AS_OF,
            approved_by="Synthetic board",
            approval_timestamp=utc_now(),
            bucket_key="pos:Synthetic snapshot",
        )
    )
    db_session.commit()
    db_session.expunge_all()
    bank = db_session.get(Bank, BANK_A)
    assert bank is not None
    for status in (KeyStatus.ACTIVE, KeyStatus.DISABLED, KeyStatus.UNAVAILABLE):
        provider.set_status(KEY_A, status)
        result = credit_concentration.monitor(db_session, CTX, bank, AS_OF)
        assert len(result.breaches) == 1
        assert result.breaches[0].key == "pos:Synthetic snapshot"
        assert not result.locked_limits


@pytest.fixture
def reference_ingestion_upload(
    db_session: Session, provider: LocalKeyProvider, tmp_path: Path
) -> IngestionBatchCreate:
    mapping = MappingConfig(
        field_mappings={
            "counterparty": EntityMapping(
                source_table="Customers",
                fields={
                    "source_reference": "CustomerId",
                    "name": "CustomerName",
                    "counterparty_type": "Type",
                },
            ),
            "position": EntityMapping(
                source_table="Loans",
                fields={
                    "source_reference": "Reference",
                    "position_type": "Type",
                    "currency": "Currency",
                    "balance": "Balance",
                    "counterparty_reference": "CustomerId",
                },
            ),
        }
    )
    ingestion.create_mapping_config(
        db_session,
        CTX,
        BANK_A,
        MappingConfigCreate(
            source_system="EXCEL_CSV",
            name="Synthetic reference lookup",
            config=mapping,
            activate=True,
            reason="Synthetic reference lookup",
        ),
    )
    workbook = Workbook()
    customers = workbook.active
    assert customers is not None
    customers.title = "Customers"
    customers.append(["CustomerId", "CustomerName", "Type"])
    customers.append([IDENTIFIER, "Synthetic replacement", "RETAIL_INDIVIDUAL"])
    loans = cast(Worksheet, workbook.create_sheet("Loans"))
    loans.append(["Reference", "Type", "Currency", "Balance", "CustomerId"])
    loans.append(["historical-loan", "LOAN", "GHS", 100, "historical-party"])
    loans.append(["replacement-loan", "LOAN", "GHS", 200, IDENTIFIER])
    path = tmp_path / "reference-lookups.xlsx"
    workbook.save(path)
    return IngestionBatchCreate(
        source_system="EXCEL_CSV",
        as_of_date=AS_OF,
        location=str(path),
        mapping_config_id=None,
        reason="Synthetic reference lookup",
    )


@pytest.mark.parametrize("connected", [False, True])
def test_ingestion_reference_lookups_leave_unused_fields_unread(
    db_session: Session,
    provider: LocalKeyProvider,
    reference_ingestion_upload: IngestionBatchCreate,
    connected: bool,
) -> None:
    if connected:
        connect(db_session, provider)
    historical = counterparty(
        db_session, reference="historical-party", as_of=date(2026, 5, 31), include_package=False
    )
    current = counterparty(db_session, include_package=False)
    prior_date = counterparty(db_session, as_of=date(2026, 5, 31), include_package=False)
    other_source = counterparty(db_session, source_system="API_PUSH", include_package=False)
    other_bank = counterparty(
        db_session, bank_id=BANK_B, reference="historical-party", include_package=False
    )
    historical_id, prior_id, current_id, other_source_id, other_bank_id = (
        historical.id,
        prior_date.id,
        current.id,
        other_source.id,
        other_bank.id,
    )
    db_session.commit()
    for record_id in (historical_id, prior_id, current_id, other_source_id, other_bank_id):
        db_session.execute(
            text(
                "UPDATE canonical_counterparties SET name = :broken, group_reference = :broken, "
                "external_identifiers = :json, attributes = :json WHERE id = :id"
            ).bindparams(bindparam("id", type_=Uuid())),
            {
                "id": record_id,
                "broken": "aqs-bank-field-v1:[]",
                "json": json.dumps({"aqs-bank-field-v1": []}),
            },
        )
    db_session.commit()
    db_session.expunge_all()
    storage = InMemoryStorageClient()
    if connected:
        for status in (KeyStatus.DISABLED, KeyStatus.UNAVAILABLE):
            provider.set_status(KEY_A, status)
            with pytest.raises(KeyUnavailableError):
                ingestion.start_ingestion(
                    db_session, CTX, BANK_A, reference_ingestion_upload, storage
                )
            db_session.rollback()
            assert list(db_session.scalars(select(CanonicalPositionSnapshot.id))) == []
        provider.set_status(KEY_A, KeyStatus.ACTIVE)
    result = ingestion.start_ingestion(db_session, CTX, BANK_A, reference_ingestion_upload, storage)
    assert result.batch.status in {"accepted", "accepted_with_warnings"}
    db_session.expunge_all()
    replacement_id = db_session.scalar(
        select(CanonicalCounterparty.id).where(
            CanonicalCounterparty.ingestion_batch_id == result.batch.id,
            CanonicalCounterparty.bank_id == BANK_A,
        )
    )
    assert replacement_id is not None
    snapshots = db_session.execute(
        select(
            CanonicalPositionSnapshot.source_reference,
            CanonicalPositionSnapshot.counterparty_id,
            CanonicalPositionSnapshot.balance,
        ).where(CanonicalPositionSnapshot.ingestion_batch_id == result.batch.id)
    ).all()
    assert set(snapshots) == {
        ("historical-loan", historical_id, Decimal(100)),
        ("replacement-loan", replacement_id, Decimal(200)),
    }
    generations: dict[UUID, UUID | None] = {
        record_id: successor_id
        for record_id, successor_id in db_session.execute(
            select(
                CanonicalCounterparty.id,
                CanonicalCounterparty.superseded_by,
            )
        ).tuples()
    }
    assert generations[current_id] == replacement_id
    assert all(
        generations[record_id] is None
        for record_id in (
            historical_id,
            prior_id,
            other_source_id,
            other_bank_id,
            replacement_id,
        )
    )
    replacement = db_session.get(CanonicalCounterparty, replacement_id)
    assert replacement is not None and replacement.name == "Synthetic replacement"
    assert replacement.source_reference == IDENTIFIER


@pytest.mark.parametrize("status", [KeyStatus.DISABLED, KeyStatus.UNAVAILABLE])
def test_numeric_saved_bi_configuration_loads_without_a_bank_key(
    db_session: Session, provider: LocalKeyProvider, status: KeyStatus
) -> None:
    connect(db_session, provider)
    query = BiPackQuery(measures=["loans.balance_rc"], window="as_of")
    spec = BiDashboardSpec.model_validate(
        {
            "widgets": [
                {
                    "id": "loans",
                    "kind": "table",
                    "title": "Loans",
                    "query": query.model_dump(mode="json"),
                }
            ],
            "layout": [{"i": "loans", "x": 0, "y": 0, "w": 6, "h": 4}],
        }
    ).model_dump(mode="json")
    version = saved_version(db_session, spec)
    subscription = BiSubscription(
        organization_id=ORG_1,
        bank_id=BANK_A,
        name="Numeric subscription",
        owner_user_id=USER_1,
        query=query.for_period(AS_OF).model_dump(mode="json"),
        artifact_format="csv",
        cadence="daily",
        hour=7,
        minute=30,
        recipient_user_ids=[],
    )
    filters = [{"member": "position.currency", "op": "eq", "values": ["GHS"]}]
    alert = BiAlert(
        organization_id=ORG_1,
        bank_id=BANK_A,
        owner_user_id=USER_1,
        name="Numeric alert",
        measure_id="loans.balance_rc",
        filters=filters,
        direction="above",
        threshold_basis="stated",
        threshold=Decimal(100),
        notify_user_ids=[],
    )
    db_session.add_all([version, subscription, alert])
    db_session.commit()
    ids = version.id, subscription.id, alert.id
    db_session.expunge_all()
    provider.set_status(KEY_A, status)
    db_session.info[SESSION_STATE] = PresentationState()
    loaded_version = db_session.get(BiDashboardVersion, ids[0])
    loaded_subscription = db_session.get(BiSubscription, ids[1])
    loaded_alert = db_session.get(BiAlert, ids[2])
    assert loaded_version is not None and loaded_version.spec == spec
    assert loaded_subscription is not None and loaded_subscription.query == query.for_period(
        AS_OF
    ).model_dump(mode="json")
    assert loaded_alert is not None and loaded_alert.filters == filters
    assert cast(PresentationState, db_session.info[SESSION_STATE]).banks == {}
    db_session.info.pop(SESSION_STATE)


@pytest.mark.parametrize("status", [KeyStatus.DISABLED, KeyStatus.UNAVAILABLE])
def test_empty_register_recovers_after_first_protected_write_refuses(
    db_client: TestClient,
    db_session: Session,
    provider: LocalKeyProvider,
    status: KeyStatus,
) -> None:
    connect(db_session, provider)
    payload: dict[str, object] = {
        "reason": "Synthetic first entry",
        "party_type": "individual",
        "full_name": NAME,
        "contact": {},
        "roles": [],
    }
    profile_url = f"/api/v1/banks/{BANK_A}/institution-profile"
    parties_url = f"/api/v1/banks/{BANK_A}/related-parties"
    provider.set_status(KEY_A, status)
    refused = db_client.post(parties_url, headers=headers(), json=payload)
    assert refused.status_code == 503, refused.text
    reason = "revoked" if status == KeyStatus.DISABLED else "unavailable"
    assert refused.json()["error"]["code"] == f"bank_key_{reason}"
    empty = db_client.get(profile_url, headers=headers())
    assert empty.status_code == 200, empty.text
    assert empty.json()["related_parties"] == []
    assert json.loads(empty.headers["X-Aequoros-Protected-Data"]) == {BANK_A: reason}

    unavailable = db_client.get(f"/api/v1/banks/{BANK_A}", headers=headers())
    assert unavailable.status_code == 200, unavailable.text
    assert json.loads(unavailable.headers["X-Aequoros-Protected-Data"]) == {BANK_A: reason}
    unauthorized = db_client.get(f"/api/v1/banks/{BANK_A}", headers=headers(org_id=ORG_2))
    assert unauthorized.status_code == 404, unauthorized.text
    assert "X-Aequoros-Protected-Data" not in unauthorized.headers

    provider.set_status(KEY_A, KeyStatus.ACTIVE)
    recovered = db_client.get(f"/api/v1/banks/{BANK_A}", headers=headers())
    assert recovered.status_code == 200, recovered.text
    assert json.loads(recovered.headers["X-Aequoros-Protected-Data"]) == {BANK_A: "available"}
    assert "no-store" in recovered.headers["cache-control"]
    still_empty = db_client.get(profile_url, headers=headers())
    assert still_empty.json()["related_parties"] == []
    saved = db_client.post(parties_url, headers=headers(), json=payload)
    assert saved.status_code == 201, saved.text
    assert saved.json()["full_name"] == NAME


def test_unconnected_bank_read_does_not_confirm_key_availability(
    db_client: TestClient, provider: LocalKeyProvider
) -> None:
    response = db_client.get(f"/api/v1/banks/{BANK_A}", headers=headers())
    assert response.status_code == 200, response.text
    assert "X-Aequoros-Protected-Data" not in response.headers


@pytest.mark.parametrize("presentation_support", [False, True])
def test_bank_recovery_read_does_not_unwrap_for_support(
    db_session: Session,
    provider: LocalKeyProvider,
    monkeypatch: pytest.MonkeyPatch,
    presentation_support: bool,
) -> None:
    connect(db_session, provider)

    def refuse_unwrap(*_args: object, **_kwargs: object) -> bytes:
        pytest.fail("Support recovery must not unwrap bank data")

    monkeypatch.setattr(provider, "unwrap", refuse_unwrap)
    presentation = PresentationState(support=presentation_support)
    db_session.info[SESSION_STATE] = presentation
    ctx = TenantContext(
        organization_id=ORG_1,
        impersonation_context="synthetic-support",
        actor_operator="synthetic-operator@example.test",
    )
    assert banks.get_bank(db_session, ctx, BANK_A).id == BANK_A
    assert presentation.banks == {BANK_A: "support"}
    db_session.info.pop(SESSION_STATE)


@pytest.mark.parametrize("status", ["deferred", "failed", "completed", None])
def test_dedup_status_remains_operational_without_unwrapping_the_report(
    db_session: Session,
    provider: LocalKeyProvider,
    monkeypatch: pytest.MonkeyPatch,
    status: str | None,
) -> None:
    connect(db_session, provider)
    batch = IngestionBatch(
        organization_id=ORG_1,
        bank_id=BANK_A,
        source_system="EXCEL_CSV",
        adapter_version="synthetic",
        extraction_mode="full",
        status="accepted",
        as_of_date=AS_OF,
        etl_report={"dedup_status": "deferred", "name": NAME},
    )
    db_session.add(batch)
    db_session.commit()
    assert batch.dedup_status == "deferred"
    batch.etl_report = {"dedup_status": status, "name": NAME} if status else None
    db_session.commit()
    batch_id = batch.id
    db_session.expunge_all()

    def refuse_unwrap(*_args: object, **_kwargs: object) -> bytes:
        pytest.fail("Operational status must not unwrap the report during support access")

    monkeypatch.setattr(provider, "unwrap", refuse_unwrap)
    db_session.info[SESSION_SUPPORT] = True
    try:
        loaded = db_session.get(IngestionBatch, batch_id)
        assert loaded is not None and loaded.dedup_status == status
        assert loaded.etl_report == {LOCKED_JSON: "support"}
    finally:
        db_session.info.pop(SESSION_SUPPORT)


@pytest.mark.parametrize("connected", [False, True])
@pytest.mark.parametrize("mutation", ["insert", "update"])
def test_plaintext_counterparty_name_bound_applies_before_orm_persistence(
    db_session: Session,
    provider: LocalKeyProvider,
    connected: bool,
    mutation: str,
) -> None:
    if connected:
        connect(db_session, provider)
    party = counterparty(db_session, name="A" * 255, include_package=False)
    party.group_reference = "G"
    party_id = party.id
    db_session.commit()
    for invalid in ("A" * 256, "A" * 300):
        with pytest.raises(ValueError, match="255"), db_session.begin_nested():
            if mutation == "insert":
                counterparty(db_session, reference="CP2", name=invalid, include_package=False)
            else:
                party.name = invalid
                db_session.flush()
        db_session.expire_all()
        loaded = db_session.get(CanonicalCounterparty, party_id)
        assert loaded is not None and loaded.name == "A" * 255
    assert (
        len(
            list(
                db_session.scalars(
                    select(CanonicalCounterparty).where(CanonicalCounterparty.bank_id == BANK_A)
                )
            )
        )
        == 1
    )


@pytest.mark.parametrize("connected", [False, True])
@pytest.mark.parametrize("bounded_field", ["name", "group_reference"])
def test_ingested_counterparty_name_bound_preserves_bi_copy_width(
    db_session: Session,
    provider: LocalKeyProvider,
    tmp_path: Path,
    connected: bool,
    bounded_field: str,
) -> None:
    if connected:
        connect(db_session, provider)
    mapping = MappingConfig(
        field_mappings={
            "counterparty": EntityMapping(
                source_table="Customers",
                fields={
                    "source_reference": "CustomerId",
                    "name": "CustomerName",
                    "counterparty_type": "Type",
                    "group_reference": "Group",
                },
            )
        }
    )
    ingestion.create_mapping_config(
        db_session,
        CTX,
        BANK_A,
        MappingConfigCreate(
            source_system="EXCEL_CSV",
            name="Synthetic bounded name",
            config=mapping,
            activate=True,
            reason="Synthetic boundary regression",
        ),
    )
    for length in (300, 255):
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Customers"
        sheet.append(["CustomerId", "CustomerName", "Type", "Group"])
        sheet.append(
            [
                "CP1",
                "A" * length if bounded_field == "name" else "A",
                "RETAIL_INDIVIDUAL",
                "G" * length if bounded_field == "group_reference" else "G",
            ]
        )
        path = tmp_path / f"synthetic-name-{length}.xlsx"
        workbook.save(path)
        result = ingestion.start_ingestion(
            db_session,
            CTX,
            BANK_A,
            IngestionBatchCreate(
                source_system="EXCEL_CSV",
                as_of_date=AS_OF,
                location=str(path),
                reason="Synthetic bounded name",
            ),
            InMemoryStorageClient(),
        )
        parties = list(
            db_session.scalars(
                select(CanonicalCounterparty).where(CanonicalCounterparty.bank_id == BANK_A)
            )
        )
        if length > 255:
            assert result.batch.records_translated == 0 and not parties
            failures = list(
                db_session.scalars(
                    select(TranslationFailure).where(
                        TranslationFailure.ingestion_batch_id == result.batch.id
                    )
                )
            )
            assert len(failures) == 1 and "255" in failures[0].error_message
        else:
            assert result.batch.status in {"accepted", "accepted_with_warnings"}
            assert len(parties) == 1
            assert (parties[0].name if bounded_field == "name" else parties[0].group_reference) == (
                "A" if bounded_field == "name" else "G"
            ) * 255
    outcome = materialize_bi_plane(db_session, organization_id=ORG_1, bank_id=BANK_A, as_of=AS_OF)
    assert outcome is not None and outcome.status == "succeeded"
    dimension = db_session.scalar(
        select(BiDimCounterparty).where(BiDimCounterparty.counterparty_id == parties[0].id)
    )
    assert dimension is not None
    assert dimension.name == (
        str(parties[0].id) if connected else ("A" * 255 if bounded_field == "name" else "A")
    )
    assert dimension.group_reference == (
        None if connected else ("G" * 255 if bounded_field == "group_reference" else "G")
    )


def assert_legacy_bi_disclosure(entity: Base, *, identity: bool) -> None:
    if isinstance(entity, BiDashboardVersion):
        value: object = entity.spec
    elif isinstance(entity, BiSubscription):
        value = entity.query
    else:
        assert isinstance(entity, BiAlert)
        value = entity.filters
    if identity:
        assert value == {LOCKED_JSON: "revoked"}
    else:
        assert value == (
            [{"borrower": NAME}] if isinstance(entity, BiAlert) else {"borrower": NAME}
        )


@pytest.mark.parametrize("refresh", [False, True])
@pytest.mark.parametrize("identity", [False, True])
def test_legacy_protected_entities_require_connected_key_on_load_and_refresh(
    db_session: Session,
    provider: LocalKeyProvider,
    refresh: bool,
    identity: bool,
) -> None:
    filters: list[dict[str, object]] = (
        [{"member": "counterparty.name", "values": [NAME]}] if identity else [{"borrower": NAME}]
    )
    spec: dict[str, object] = {"filters": filters} if identity else {"borrower": NAME}
    counterparty(db_session)
    party = RelatedParty(
        organization_id=ORG_1,
        bank_id=BANK_A,
        party_type="individual",
        full_name=NAME,
        contact={"national_id": IDENTIFIER},
    )
    db_session.add(party)
    db_session.flush()
    db_session.add(
        RelatedPartyRole(
            organization_id=ORG_1,
            party_id=party.id,
            role="director",
            other_responsibilities=NAME,
        )
    )
    db_session.add(saved_version(db_session, spec))
    db_session.add(
        BiSubscription(
            organization_id=ORG_1,
            bank_id=BANK_A,
            name="Synthetic legacy subscription",
            owner_user_id=USER_1,
            query=spec,
            artifact_format="csv",
            cadence="daily",
            hour=7,
            minute=30,
            recipient_user_ids=[],
        )
    )
    db_session.add(
        BiAlert(
            organization_id=ORG_1,
            bank_id=BANK_A,
            owner_user_id=USER_1,
            name="Synthetic legacy alert",
            measure_id="loans.balance_rc",
            filters=filters,
            direction="above",
            threshold_basis="stated",
            threshold=Decimal(100),
            notify_user_ids=[],
        )
    )
    db_session.commit()
    models: tuple[type[Base], ...] = (
        CanonicalCounterparty,
        RelatedParty,
        RelatedPartyRole,
        IngestionBatch,
        LineageRecord,
        TranslationFailure,
        RegulatoryPackage,
        BiDashboardVersion,
        BiSubscription,
        BiAlert,
    )
    connect(db_session, provider)
    stored = raw_fields(db_session)
    immutable = cast(
        object, db_session.execute(text("SELECT spec FROM bi_dashboard_versions")).scalar_one()
    )
    for model in models:
        db_session.expunge_all()
        provider.set_status(KEY_A, KeyStatus.ACTIVE)
        entity: Base | None = None
        if refresh:
            entity = db_session.scalar(select(model))
            assert entity is not None
        db_session.commit()
        provider.set_status(KEY_A, KeyStatus.DISABLED)
        presentation_read = model in {
            RegulatoryPackage,
            BiDashboardVersion,
            BiSubscription,
            BiAlert,
        }
        if presentation_read:
            if refresh:
                assert entity is not None
                db_session.refresh(entity)
            else:
                entity = db_session.scalar(select(model))
            assert entity is not None
            if isinstance(entity, RegulatoryPackage):
                assert isinstance(entity.snapshot, dict)
            else:
                assert_legacy_bi_disclosure(entity, identity=identity)
        else:
            with pytest.raises(KeyUnavailableError):
                if refresh:
                    assert entity is not None
                    db_session.refresh(entity)
                    if isinstance(entity, CanonicalCounterparty):
                        _ = entity.name
                else:
                    entity = db_session.scalar(select(model).options(undefer("*")))
        db_session.rollback()
    assert raw_fields(db_session) == stored
    assert (
        db_session.execute(text("SELECT spec FROM bi_dashboard_versions")).scalar_one() == immutable
    )
