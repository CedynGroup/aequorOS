"""The ICAAP working copy (python-docx).

The DOCX exists because a bank's reviewers redline in Word. It is an *editable*
copy, so it can never be the filed instrument — the filing is the signed PDF —
and the header on every page says exactly that: ``DRAFT — not approved ·
WORKING COPY — not the filed document``, with the rehearsal line added when the
cycle is a rehearsal (D-029). Nothing here is ever signed, stored as a filing
artifact, or counted towards an obligation.

Two mechanics worth knowing:

* **Determinism.** python-docx stamps every zip entry with wall-clock time at
  save, the same problem openpyxl has. The container is therefore re-written
  through the exporters' shared normaliser
  (``regulatory_reporting/exports/xlsx.normalize_zip``) with fixed entry
  timestamps and a pinned ``dcterms:modified``, and the core properties are set
  explicitly from the caller's ``generated_at``. Two renders of the same
  document produce the same bytes. (Cross-machine byte identity is not
  claimed: zlib output can differ by platform and version — P1-DESIGN §4.3.)
* **Unicode.** Unlike the draft PDF, which is limited to what the standard
  font can draw (DV-004), OOXML carries UTF-8: the cedi sign and Ghanaian
  vowels print exactly as typed. Only the C0 control characters XML forbids are
  made safe, by the same helper the other exporters use.

Lists are written with explicit ``•`` / ``1.`` prefixes in the ``List
Paragraph`` style rather than Word numbering definitions, because python-docx
cannot restart a numbering definition per list deterministically — a second
list in the document would continue the first one's numbers.
"""

from __future__ import annotations

import io
from collections.abc import Iterable, Sequence
from typing import Any

from docx import Document as new_document
from docx.document import Document as DocxDocument
from docx.shared import Pt
from docx.text.paragraph import Paragraph as DocxParagraph

from app.domain.icaap.document import (
    DRAFT_PROVENANCE_NOTICE,
    EXPOSURE_DRAFT_NOTICE,
    NOT_AVAILABLE,
    WORKING_COPY_NOTICE,
    BlockRender,
    DataBlockRef,
    IcaapDocument,
    ListBlock,
    ParagraphBlock,
    QuoteBlock,
    RenderBlock,
    Run,
    SectionRender,
    TableSpec,
)
from app.services.icaap.render import format as fmt
from app.services.regulatory_reporting.exports.text import xml_safe
from app.services.regulatory_reporting.exports.xlsx import normalize_zip

#: Styles the python-docx default template ships. Nothing else is referenced,
#: so the document opens with its styles intact in Word, Pages and LibreOffice.
_HEADING_STYLES = {"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3", "h4": "Heading 4"}
_BODY_STYLE = "Normal"
_QUOTE_STYLE = "Quote"
_LIST_STYLE = "List Paragraph"
_TABLE_STYLE = "Table Grid"
_LIST_INDENT_PT = 18
_NOTE_SIZE = Pt(8)
_AUTHOR = "AequorOS"


def _safe(text: str) -> str:
    """Text OOXML can carry. python-docx escapes markup itself, so the only
    hazard is the C0 control set, handled by the shared exporter helper."""
    return xml_safe(text)


class _DocxRenderer:
    def __init__(self, document: IcaapDocument) -> None:
        self._doc = document
        self._out: DocxDocument = new_document()

    # -- primitives -------------------------------------------------------
    def _paragraph(self, text: str = "", *, style: str = _BODY_STYLE) -> DocxParagraph:
        paragraph = self._out.add_paragraph(style=style)
        if text:
            paragraph.add_run(_safe(text))
        return paragraph

    def _note(self, text: str) -> None:
        if not text:
            return
        paragraph = self._out.add_paragraph(style=_BODY_STYLE)
        run = paragraph.add_run(_safe(text))
        run.italic = True
        run.font.size = _NOTE_SIZE

    def _heading(self, text: str, level: str) -> None:
        self._paragraph(text, style=_HEADING_STYLES[level])

    # -- editor IR --------------------------------------------------------
    def _render_blocks(self, blocks: Iterable[RenderBlock], *, depth: int = 0) -> None:
        for block in blocks:
            self._render_block(block, depth=depth)

    def _render_block(self, block: RenderBlock, *, depth: int = 0) -> None:
        if isinstance(block, ParagraphBlock):
            style = _HEADING_STYLES.get(block.style, _BODY_STYLE)
            self._write_runs(self._out.add_paragraph(style=style), block.runs)
            return
        if isinstance(block, ListBlock):
            self._render_list(block, depth=depth)
            return
        if isinstance(block, QuoteBlock):
            self._render_quote(block, depth=depth)
            return
        if isinstance(block, DataBlockRef):
            self._render_data_block(block.block_id)

    def _write_runs(self, paragraph: DocxParagraph, runs: Sequence[Run]) -> None:
        for run in runs:
            if run.line_break:
                paragraph.add_run().add_break()
                continue
            fact = run.fact
            # A factRef prints the figure the block was bound to, exactly as
            # the PDF prints it, so the two documents never disagree.
            text = self._doc.fact(fact.block_id, fact.fact_key) if fact is not None else run.text
            if not text:
                continue
            written = paragraph.add_run(_safe(text))
            written.bold = run.bold
            written.italic = run.italic
            written.underline = run.underline

    def _render_list(self, block: ListBlock, *, depth: int) -> None:
        number = max(1, int(block.start))
        for item in block.items:
            prefix = "• " if not block.ordered else f"{number}. "
            number += 1
            first = True
            for child in item:
                if first and isinstance(child, ParagraphBlock):
                    paragraph = self._out.add_paragraph(style=_LIST_STYLE)
                    paragraph.paragraph_format.left_indent = Pt(_LIST_INDENT_PT * (depth + 1))
                    paragraph.add_run(prefix)
                    self._write_runs(paragraph, child.runs)
                    first = False
                    continue
                self._render_block(child, depth=depth + 1)
            if first:
                # An item with no leading paragraph still needs its marker, or
                # the list silently loses an entry.
                paragraph = self._out.add_paragraph(style=_LIST_STYLE)
                paragraph.paragraph_format.left_indent = Pt(_LIST_INDENT_PT * (depth + 1))
                paragraph.add_run(prefix)

    def _render_quote(self, block: QuoteBlock, *, depth: int) -> None:
        for child in block.children:
            if isinstance(child, ParagraphBlock):
                self._write_runs(self._out.add_paragraph(style=_QUOTE_STYLE), child.runs)
                continue
            self._render_block(child, depth=depth)

    # -- data blocks ------------------------------------------------------
    def _render_data_block(self, block_id: str) -> None:
        block = self._doc.block(block_id)
        if block is None:
            self._note("A figure table referenced here is no longer part of this cycle.")
            return
        self._heading(block.title, "h3")
        meta = self._block_meta_line(block)
        if meta:
            self._note(meta)
        if block.status_note:
            self._note(block.status_note)
        if not block.tables:
            self._paragraph(NOT_AVAILABLE)
        for table in block.tables:
            self._table(table)
        for note in block.notes:
            self._note(note)

    def _block_meta_line(self, block: BlockRender) -> str:
        parts = [part for part in (block.source_label,) if part]
        if block.as_of is not None:
            parts.append(f"As at {fmt.format_date(block.as_of)}")
        parts.append(fmt.block_status_label(block.status))
        return " · ".join(parts)

    def _table(self, spec: TableSpec) -> None:
        if spec.title:
            self._heading(spec.title, "h4")
        if not spec.columns:
            self._paragraph(NOT_AVAILABLE)
            return
        table = self._out.add_table(rows=1, cols=len(spec.columns), style=_TABLE_STYLE)
        for cell, column in zip(table.rows[0].cells, spec.columns, strict=True):
            run = cell.paragraphs[0].add_run(_safe(column.label))
            run.bold = True
        for row in spec.rows:
            cells = table.add_row().cells
            for cell, column in zip(cells, spec.columns, strict=True):
                run = cell.paragraphs[0].add_run(_safe(row.cells.get(column.key, "")))
                run.bold = bool(row.emphasis)
        if not spec.rows:
            self._note("No rows recorded.")

    # -- document parts ---------------------------------------------------
    def _cover(self) -> None:
        doc = self._doc
        self._heading(doc.cycle_title, "h1")
        for line in self._banner_lines():
            self._paragraph(line, style=_BODY_STYLE).runs[0].bold = True
        if doc.is_exposure_draft:
            self._note(EXPOSURE_DRAFT_NOTICE)
        if doc.framework_digest_note:
            self._note(doc.framework_digest_note)
        rows = (
            ("Institution", doc.institution_name),
            ("Financial year", f"FY{doc.fiscal_year}"),
            ("Reporting date", fmt.format_date(doc.as_of)),
            ("Basis", doc.basis_label),
            ("Cycle", doc.cycle_kind_label),
            ("Framework", f"{doc.framework_title} {doc.framework_version}"),
            ("Reporting currency", doc.currency),
            ("Generated at", doc.generated_at.isoformat(timespec="seconds")),
        )
        table = self._out.add_table(rows=0, cols=2, style=_TABLE_STYLE)
        for label, value in rows:
            cells = table.add_row().cells
            cells[0].paragraphs[0].add_run(_safe(label)).bold = True
            cells[1].paragraphs[0].add_run(_safe(value))
        if doc.readiness_summary:
            self._heading("Readiness", "h2")
            for line in doc.readiness_summary:
                self._paragraph(line)

    def _banner_lines(self) -> tuple[str, ...]:
        return (*self._doc.watermark_lines, WORKING_COPY_NOTICE)

    def _contents(self) -> None:
        self._heading("Contents", "h1")
        if not self._doc.sections:
            self._paragraph("This cycle has no sections.")
            return
        table = self._out.add_table(rows=0, cols=3, style=_TABLE_STYLE)
        for section in self._doc.sections:
            cells = table.add_row().cells
            cells[0].paragraphs[0].add_run(_safe(f"({section.letter})"))
            cells[1].paragraphs[0].add_run(_safe(section.title))
            cells[2].paragraphs[0].add_run(_safe(section.content_label))

    def _section(self, section: SectionRender) -> None:
        self._heading(f"({section.letter}) {section.title}", "h1")
        meta = " · ".join(part for part in (section.citation_label, section.content_label) if part)
        if meta:
            self._note(meta)
        if section.source_status and section.source_status != "sourced":
            self._note(fmt.pending_primary_text_note(self._doc.regulator_short))
        if section.blocks:
            self._render_blocks(section.blocks)
        elif section.empty_note:
            self._note(section.empty_note)

    def _checklist(self) -> None:
        self._heading("Requirement checklist", "h1")
        lines = [(section, line) for section in self._doc.sections for line in section.checklist]
        if not lines:
            self._paragraph("No requirements are recorded for this framework.")
            return
        table = self._out.add_table(rows=1, cols=5, style=_TABLE_STYLE)
        headers = ("Section", "Requirement", "Citation", "Status", "Reason")
        for cell, header in zip(table.rows[0].cells, headers, strict=True):
            cell.paragraphs[0].add_run(_safe(header)).bold = True
        for section, line in lines:
            cells = table.add_row().cells
            cells[0].paragraphs[0].add_run(_safe(f"({section.letter})"))
            cells[1].paragraphs[0].add_run(_safe(line.text))
            cells[2].paragraphs[0].add_run(_safe(line.citation_label))
            cells[3].paragraphs[0].add_run(_safe(fmt.requirement_status_label(line.status)))
            # REG-ICAAP-008: the reason travels with the document.
            cells[4].paragraphs[0].add_run(_safe(line.reason or ""))

    def _attachments(self) -> None:
        self._heading("Attachments", "h1")
        if not self._doc.attachments:
            self._paragraph("No documents have been attached to this cycle.")
            return
        table = self._out.add_table(rows=1, cols=5, style=_TABLE_STYLE)
        headers = ("Kind", "Title", "Checksum", "Uploaded", "Status")
        for cell, header in zip(table.rows[0].cells, headers, strict=True):
            cell.paragraphs[0].add_run(_safe(header)).bold = True
        for attachment in self._doc.attachments:
            cells = table.add_row().cells
            cells[0].paragraphs[0].add_run(_safe(attachment.kind_title))
            cells[1].paragraphs[0].add_run(_safe(attachment.title))
            cells[2].paragraphs[0].add_run(_safe(attachment.sha256_prefix))
            cells[3].paragraphs[0].add_run(_safe(fmt.format_date(attachment.uploaded_on)))
            cells[4].paragraphs[0].add_run(_safe("Withdrawn" if attachment.withdrawn else "Active"))

    def _provenance(self) -> None:
        doc = self._doc
        self._heading("Provenance", "h1")
        lines = [
            f"Framework: {doc.framework_title} {doc.framework_version} ({doc.framework_code})",
            f"Framework status: {fmt.humanize(doc.framework_status)}",
            f"Framework digest: {doc.framework_digest}",
            f"Generated at: {doc.generated_at.isoformat(timespec='seconds')}",
            DRAFT_PROVENANCE_NOTICE,
            WORKING_COPY_NOTICE,
        ]
        if doc.framework_digest_note:
            lines.insert(3, doc.framework_digest_note)
        if doc.is_exposure_draft:
            lines.insert(2, EXPOSURE_DRAFT_NOTICE)
        for line in lines:
            self._paragraph(line)
        self._heading("Figure sources", "h2")
        sourced = [block for block in doc.blocks.values() if block.provenance or block.source_label]
        if not sourced:
            self._paragraph("No figure tables are bound in this cycle.")
            return
        for block in sourced:
            self._paragraph(block.title).runs[0].bold = True
            detail = [f"Status: {fmt.block_status_label(block.status)}"]
            if block.source_label:
                detail.append(block.source_label)
            if block.as_of is not None:
                detail.append(f"As at {fmt.format_date(block.as_of)}")
            detail.extend(f"{label}: {value}" for label, value in block.provenance)
            self._paragraph(" · ".join(detail))

    # -- assembly ---------------------------------------------------------
    def _page_header(self) -> None:
        """The notice on every Word section, i.e. on every page."""
        text = " · ".join(self._banner_lines())
        for section in self._out.sections:
            paragraph = section.header.paragraphs[0]
            paragraph.text = ""
            paragraph.add_run(_safe(text)).bold = True

    def _core_properties(self) -> None:
        doc = self._doc
        pinned = doc.generated_at.replace(tzinfo=None, microsecond=0)
        props: Any = self._out.core_properties
        props.author = _AUTHOR
        props.last_modified_by = _AUTHOR
        props.title = doc.cycle_title
        props.subject = f"ICAAP FY{doc.fiscal_year} ({doc.basis_label})"
        props.keywords = "ICAAP working copy"
        props.comments = WORKING_COPY_NOTICE
        props.category = doc.framework_code
        props.created = pinned
        props.modified = pinned
        props.last_printed = pinned
        props.revision = 1

    def build(self) -> bytes:
        self._cover()
        self._contents()
        for section in self._doc.sections:
            self._section(section)
        self._checklist()
        self._attachments()
        self._provenance()
        self._page_header()
        self._core_properties()
        buffer = io.BytesIO()
        self._out.save(buffer)
        return normalize_zip(
            buffer.getvalue(), self._doc.generated_at.replace(tzinfo=None, microsecond=0)
        )


def render_docx(document: IcaapDocument) -> bytes:
    """One ICAAP document as DOCX bytes. Same input, same bytes."""
    return _DocxRenderer(document).build()


__all__ = ["render_docx"]
